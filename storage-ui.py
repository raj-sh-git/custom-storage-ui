import os
import io
import re
import math
import base64
import json
import uuid
import mimetypes
import csv
import openpyxl
import zipfile
from datetime import datetime, timezone, timedelta
from flask import (
    Flask, render_template, request, redirect, url_for, session,
    send_file, flash, Blueprint, jsonify, Response, abort, has_request_context
)
from flask_session import Session
from werkzeug.utils import secure_filename
from werkzeug.middleware.proxy_fix import ProxyFix
from werkzeug.security import generate_password_hash, check_password_hash
from azure.storage.blob import BlobServiceClient, ContentSettings
from azure.storage.fileshare import ShareServiceClient
from azure.storage.queue import QueueServiceClient
from azure.data.tables import TableServiceClient, TableEntity, UpdateMode
from azure.core.exceptions import ResourceNotFoundError, HttpResponseError
from azure.identity import DefaultAzureCredential, ClientSecretCredential

# -----------------------
# App Configuration
# -----------------------
app = Flask(__name__, static_url_path="/storage-ui/static")
app.wsgi_app = ProxyFix(app.wsgi_app, x_for=1, x_proto=1, x_host=1, x_prefix=1)

# Server-side Session Configuration
SESSION_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "flask_sessions")
os.makedirs(SESSION_DIR, exist_ok=True)

app.secret_key = os.environ.get('SECRET_KEY', str(uuid.uuid4()))
app.config.update(
    SESSION_TYPE='filesystem',
    SESSION_FILE_DIR=SESSION_DIR,
    SESSION_PERMANENT=True,
    PERMANENT_SESSION_LIFETIME=86400, # 24 hours
)
Session(app)

# System Table Names & Security Policies (configurable via environment variables)
USER_TABLE = os.environ.get('USER_TABLE', 'StorageUIUsers')
LOGS_TABLE = os.environ.get('LOGS_TABLE', 'StorageUIActivityLogs')
DEFAULT_PASSWORD_EXPIRY_DAYS = int(os.environ.get('DEFAULT_PASSWORD_EXPIRY_DAYS', '90'))
ACTIVITY_LOG_RETENTION_DAYS = int(os.environ.get('ACTIVITY_LOG_RETENTION_DAYS', '30'))

@app.context_processor
def inject_user_context():
    user = session.get('user') if (has_request_context() and 'user' in session) else None
    return {
        'current_user': user,
        'user_role': (user.get('role') if user else 'contributor'),
        'is_admin': (user.get('role') == 'admin') if user else False,
        'is_reader': (user.get('role') == 'reader') if user else False
    }

# ---------- Blueprint ----------
ui = Blueprint("ui", __name__, url_prefix="/storage-ui")

# -----------------------
# Helpers & Client Providers
# -----------------------
def parse_account_name(conn_str):
    for part in conn_str.split(";"):
        if part.strip().startswith("AccountName="):
            return part.split("=", 1)[1]
    return "Storage Account"

def auto_connect_from_env_if_available():
    """Auto-connects to Azure Storage if environment variables are set."""
    if 'auth_method' in session and 'account_name' in session:
        return True
        
    conn_str = os.environ.get('AZURE_STORAGE_CONNECTION_STRING')
    if conn_str:
        account_name = parse_account_name(conn_str)
        session['auth_method'] = 'Connection String'
        session['account_name'] = account_name
        session['conn_string'] = conn_str
        session['is_env_managed'] = True
        return True
        
    account_name = os.environ.get('AZURE_STORAGE_ACCOUNT_NAME')
    if account_name:
        account_key = os.environ.get('AZURE_STORAGE_ACCOUNT_KEY')
        if account_key:
            session['auth_method'] = 'Access Key'
            session['account_name'] = account_name
            session['account_key'] = account_key
            session['is_env_managed'] = True
            return True
        else:
            session['auth_method'] = 'DefaultAzureCredential'
            session['account_name'] = account_name
            session['is_env_managed'] = True
            return True
            
    return False

def is_storage_connected():
    if has_request_context():
        if 'auth_method' in session and 'account_name' in session:
            return True
        return auto_connect_from_env_if_available()
    return True

def require_auth():
    """Requires both storage connection and valid user login session."""
    if not is_storage_connected():
        return False
    user = session.get('user')
    return bool(user and user.get('is_authenticated'))

def is_admin():
    return require_auth() and session.get('user', {}).get('role') == 'admin'

def has_write_permission():
    return require_auth() and session.get('user', {}).get('role') in ['admin', 'contributor']

def get_blob_service():
    if not is_storage_connected():
        raise ValueError("Not authenticated with Azure Storage")
    method = session.get('auth_method')
    if method == 'Connection String':
        return BlobServiceClient.from_connection_string(session['conn_string'])
    elif method == 'Access Key':
        conn_str = f"DefaultEndpointsProtocol=https;AccountName={session['account_name']};AccountKey={session['account_key']};EndpointSuffix=core.windows.net"
        return BlobServiceClient.from_connection_string(conn_str)
    elif method == 'Service Principal':
        credential = ClientSecretCredential(session['tenant_id'], session['client_id'], session['client_secret'])
        return BlobServiceClient(account_url=f"https://{session['account_name']}.blob.core.windows.net", credential=credential)
    elif method == 'DefaultAzureCredential':
        credential = DefaultAzureCredential()
        return BlobServiceClient(account_url=f"https://{session['account_name']}.blob.core.windows.net", credential=credential)
    raise ValueError("Not authenticated with Azure Storage")

def get_share_service():
    if not is_storage_connected():
        raise ValueError("Not authenticated with Azure Storage")
    method = session.get('auth_method')
    if method == 'Connection String':
        return ShareServiceClient.from_connection_string(session['conn_string'])
    elif method == 'Access Key':
        conn_str = f"DefaultEndpointsProtocol=https;AccountName={session['account_name']};AccountKey={session['account_key']};EndpointSuffix=core.windows.net"
        return ShareServiceClient.from_connection_string(conn_str)
    elif method == 'Service Principal':
        credential = ClientSecretCredential(session['tenant_id'], session['client_id'], session['client_secret'])
        return ShareServiceClient(account_url=f"https://{session['account_name']}.file.core.windows.net", credential=credential)
    elif method == 'DefaultAzureCredential':
        credential = DefaultAzureCredential()
        return ShareServiceClient(account_url=f"https://{session['account_name']}.file.core.windows.net", credential=credential)
    raise ValueError("Not authenticated with Azure Storage")

def get_queue_service():
    if not is_storage_connected():
        raise ValueError("Not authenticated with Azure Storage")
    method = session.get('auth_method')
    if method == 'Connection String':
        return QueueServiceClient.from_connection_string(session['conn_string'])
    elif method == 'Access Key':
        conn_str = f"DefaultEndpointsProtocol=https;AccountName={session['account_name']};AccountKey={session['account_key']};EndpointSuffix=core.windows.net"
        return QueueServiceClient.from_connection_string(conn_str)
    elif method == 'Service Principal':
        credential = ClientSecretCredential(session['tenant_id'], session['client_id'], session['client_secret'])
        return QueueServiceClient(account_url=f"https://{session['account_name']}.queue.core.windows.net", credential=credential)
    elif method == 'DefaultAzureCredential':
        credential = DefaultAzureCredential()
        return QueueServiceClient(account_url=f"https://{session['account_name']}.queue.core.windows.net", credential=credential)
    raise ValueError("Not authenticated with Azure Storage")

def get_table_service():
    if not is_storage_connected():
        raise ValueError("Not authenticated with Azure Storage")
    method = session.get('auth_method')
    if method == 'Connection String':
        return TableServiceClient.from_connection_string(session['conn_string'])
    elif method == 'Access Key':
        conn_str = f"DefaultEndpointsProtocol=https;AccountName={session['account_name']};AccountKey={session['account_key']};EndpointSuffix=core.windows.net"
        return TableServiceClient.from_connection_string(conn_str)
    elif method == 'Service Principal':
        credential = ClientSecretCredential(session['tenant_id'], session['client_id'], session['client_secret'])
        return TableServiceClient(endpoint=f"https://{session['account_name']}.table.core.windows.net", credential=credential)
    elif method == 'DefaultAzureCredential':
        credential = DefaultAzureCredential()
        return TableServiceClient(endpoint=f"https://{session['account_name']}.table.core.windows.net", credential=credential)
    raise ValueError("Not authenticated with Azure Storage")

# -----------------------
# System Tables & User Management Helpers
# -----------------------
def ensure_system_tables_exist():
    """Ensures USER_TABLE and LOGS_TABLE exist in Azure Table Storage."""
    try:
        table_svc = get_table_service()
        try:
            table_svc.create_table_if_not_exists(USER_TABLE)
        except Exception as e:
            print(f"Notice: USER_TABLE check: {e}")
        try:
            table_svc.create_table_if_not_exists(LOGS_TABLE)
        except Exception as e:
            print(f"Notice: LOGS_TABLE check: {e}")
    except Exception as e:
        print(f"Warning: could not verify system tables: {e}")

def check_has_any_users():
    """Checks whether any user accounts exist in USER_TABLE."""
    try:
        table_svc = get_table_service()
        table_client = table_svc.get_table_client(USER_TABLE)
        entities = list(table_client.list_entities(results_per_page=1))
        return len(entities) > 0
    except Exception as e:
        print(f"Error checking users in {USER_TABLE}: {e}")
        return False

def get_user_by_username(username):
    if not username:
        return None
    try:
        table_svc = get_table_service()
        table_client = table_svc.get_table_client(USER_TABLE)
        return dict(table_client.get_entity(partition_key='user', row_key=username.lower().strip()))
    except (ResourceNotFoundError, HttpResponseError):
        return None
    except Exception as e:
        print(f"Error fetching user {username}: {e}")
        return None

def save_user(username, password=None, password_hash=None, email=None, display_name=None, role=None, is_active=None, must_change_password=None, password_expiry_days=None, update_login=False):
    table_svc = get_table_service()
    table_client = table_svc.get_table_client(USER_TABLE)
    row_key = username.lower().strip()

    existing = None
    try:
        existing = dict(table_client.get_entity(partition_key='user', row_key=row_key))
    except (ResourceNotFoundError, HttpResponseError):
        existing = None

    entity = {
        'PartitionKey': 'user',
        'RowKey': row_key,
        'username': username.strip(),
    }

    now_iso = datetime.now(timezone.utc).isoformat()

    if existing is None:
        entity['created_at'] = now_iso
        entity['email'] = (email or '').strip()
        entity['display_name'] = (display_name or username).strip()
        entity['role'] = (role or 'contributor').lower().strip()
        entity['is_active'] = True if is_active is None else bool(is_active)
        entity['must_change_password'] = False if must_change_password is None else bool(must_change_password)
        entity['password_expiry_days'] = int(password_expiry_days if password_expiry_days is not None else DEFAULT_PASSWORD_EXPIRY_DAYS)
        entity['password_last_set_at'] = now_iso
        entity['last_login'] = ''
    else:
        if email is not None:
            entity['email'] = email.strip()
        if display_name is not None:
            entity['display_name'] = display_name.strip()
        if role is not None:
            entity['role'] = role.lower().strip()
        if is_active is not None:
            entity['is_active'] = bool(is_active)
        if must_change_password is not None:
            entity['must_change_password'] = bool(must_change_password)
        if password_expiry_days is not None:
            try:
                entity['password_expiry_days'] = int(password_expiry_days)
            except (ValueError, TypeError):
                entity['password_expiry_days'] = DEFAULT_PASSWORD_EXPIRY_DAYS

    if password:
        entity['password_hash'] = generate_password_hash(password, method='scrypt')
        entity['password_last_set_at'] = now_iso
    elif password_hash:
        entity['password_hash'] = password_hash
        entity['password_last_set_at'] = now_iso

    if update_login:
        entity['last_login'] = now_iso

    table_client.upsert_entity(entity=entity, mode=UpdateMode.MERGE)
    return entity

def get_all_users():
    try:
        table_svc = get_table_service()
        table_client = table_svc.get_table_client(USER_TABLE)
        entities = list(table_client.query_entities("PartitionKey eq 'user'"))
        users = []
        for e in entities:
            u = dict(e)
            u.setdefault('display_name', u.get('username', ''))
            u.setdefault('email', '')
            u.setdefault('role', 'contributor')
            u.setdefault('is_active', True)
            u.setdefault('must_change_password', False)
            u.setdefault('password_expiry_days', DEFAULT_PASSWORD_EXPIRY_DAYS)
            u.setdefault('password_last_set_at', u.get('created_at', ''))
            u.setdefault('last_login', '')
            u.setdefault('created_at', '')

            # Calculate password age and expiration status
            try:
                expiry_days = int(u.get('password_expiry_days', DEFAULT_PASSWORD_EXPIRY_DAYS) or 0)
            except Exception:
                expiry_days = DEFAULT_PASSWORD_EXPIRY_DAYS
            u['password_expiry_days'] = expiry_days

            last_set = u.get('password_last_set_at') or u.get('created_at')
            if expiry_days <= 0:
                u['expiry_status'] = 'Never Expires'
                u['days_remaining'] = None
                u['is_expired'] = False
            elif last_set:
                try:
                    last_set_dt = datetime.fromisoformat(last_set.replace('Z', '+00:00'))
                    age_days = (datetime.now(timezone.utc) - last_set_dt).days
                    rem_days = expiry_days - age_days
                    u['days_remaining'] = rem_days
                    u['password_age_days'] = age_days
                    if rem_days <= 0:
                        u['expiry_status'] = 'Expired'
                        u['is_expired'] = True
                    else:
                        u['expiry_status'] = f"{rem_days} days left"
                        u['is_expired'] = False
                except Exception:
                    u['expiry_status'] = f"{expiry_days} days"
                    u['days_remaining'] = expiry_days
                    u['is_expired'] = False
            else:
                u['expiry_status'] = f"{expiry_days} days"
                u['days_remaining'] = expiry_days
                u['is_expired'] = False

            users.append(u)
        users.sort(key=lambda x: x.get('created_at', '') or x.get('username', ''))
        return users
    except Exception as e:
        print(f"Error listing users: {e}")
        return []

def delete_user_by_username(username):
    table_svc = get_table_service()
    table_client = table_svc.get_table_client(USER_TABLE)
    table_client.delete_entity(partition_key='user', row_key=username.lower().strip())

def parse_users_file(file_obj, filename):
    """Parses a CSV or Excel file and returns structured user entries for preview or batch import."""
    rows = []
    errors = []
    fn = filename.lower()
    
    try:
        if fn.endswith('.csv'):
            content = file_obj.read().decode('utf-8-sig', errors='ignore')
            reader = csv.DictReader(io.StringIO(content))
            for r in reader:
                rows.append(r)
        elif fn.endswith('.xlsx'):
            wb = openpyxl.load_workbook(file_obj, data_only=True)
            ws = wb.active
            headers = None
            for row_idx, row in enumerate(ws.iter_rows(values_only=True)):
                if row_idx == 0:
                    headers = [str(h).strip().lower() if h is not None else '' for h in row]
                else:
                    if not any(row):
                        continue
                    row_dict = {}
                    for i, val in enumerate(row):
                        if headers and i < len(headers) and headers[i]:
                            row_dict[headers[i]] = str(val).strip() if val is not None else ''
                    rows.append(row_dict)
        else:
            return [], {'total': 0, 'valid': 0, 'invalid': 0, 'existing': 0}, ["Unsupported file format. Please upload a .csv or .xlsx file."]
    except Exception as e:
        return [], {'total': 0, 'valid': 0, 'invalid': 0, 'existing': 0}, [f"Error reading file: {e}"]

    parsed_users = []
    valid_count = 0
    invalid_count = 0
    existing_count = 0

    for idx, row in enumerate(rows, start=2):
        n_row = {str(k).strip().lower(): str(v).strip() for k, v in row.items() if k is not None}
        username = n_row.get('username', '').strip()
        email = n_row.get('email', '').strip()
        display_name = n_row.get('display_name', n_row.get('displayname', n_row.get('name', username))).strip() or username
        password = n_row.get('password', '').strip()
        role = n_row.get('role', 'contributor').strip().lower()
        if role not in ['admin', 'contributor', 'reader']:
            role = 'contributor'
            
        enforce_reset_val = n_row.get('enforcepasswordreset', n_row.get('enforce_reset', n_row.get('must_change_password', 'yes'))).strip().lower()
        must_change = enforce_reset_val in ['yes', 'true', '1', 'y']
        
        raw_expiry = n_row.get('password_expiry_days', n_row.get('expiry_days', n_row.get('expiry', '')))
        try:
            password_expiry_days = int(raw_expiry) if raw_expiry else DEFAULT_PASSWORD_EXPIRY_DAYS
        except Exception:
            password_expiry_days = DEFAULT_PASSWORD_EXPIRY_DAYS

        row_errors = []
        if not username:
            row_errors.append("Missing username")
        if not password:
            row_errors.append(f"Missing password for '{username or 'user'}'")
            
        is_valid = len(row_errors) == 0
        already_exists = False
        if username:
            try:
                if get_user_by_username(username) is not None:
                    already_exists = True
                    existing_count += 1
            except Exception:
                pass

        if is_valid:
            valid_count += 1
        else:
            invalid_count += 1
            errors.append(f"Row {idx}: {', '.join(row_errors)}")

        user_item = {
            'row_index': idx,
            'username': username,
            'display_name': display_name,
            'email': email,
            'password': password,
            'masked_password': '••••••••' if password else '(empty)',
            'role': role,
            'must_change_password': must_change,
            'password_expiry_days': password_expiry_days,
            'is_valid': is_valid,
            'already_exists': already_exists,
            'error_msg': ', '.join(row_errors) if row_errors else ''
        }
        parsed_users.append(user_item)

    summary = {
        'total': len(parsed_users),
        'valid': valid_count,
        'invalid': invalid_count,
        'existing': existing_count
    }
    return parsed_users, summary, errors

def bulk_create_users_from_file(file_obj, filename):
    parsed_users, summary, errors = parse_users_file(file_obj, filename)
    created = 0
    skipped = 0
    
    for u in parsed_users:
        if not u['is_valid']:
            skipped += 1
            continue
        try:
            save_user(
                username=u['username'],
                password=u['password'],
                email=u['email'],
                display_name=u['display_name'],
                role=u['role'],
                is_active=True,
                must_change_password=u['must_change_password'],
                password_expiry_days=u['password_expiry_days']
            )
            created += 1
        except Exception as e:
            skipped += 1
            errors.append(f"Row {u['row_index']} ({u['username']}): {e}")

    return created, skipped, errors

# -----------------------
# Activity Audit Logging Helper
# -----------------------
def log_activity(service, action, target, status='SUCCESS', details='', username=None, role=None):
    try:
        if not is_storage_connected():
            return
        table_svc = get_table_service()
        table_client = table_svc.get_table_client(LOGS_TABLE)

        now = datetime.now(timezone.utc)
        inv_ts = f"{9999999999 - int(now.timestamp()):010d}"
        row_key = f"{inv_ts}_{uuid.uuid4().hex[:8]}"

        user_info = {}
        ip = '127.0.0.1'
        if has_request_context():
            try:
                user_info = session.get('user', {}) if session else {}
            except Exception:
                user_info = {}
            try:
                ip = request.headers.get('X-Forwarded-For', request.remote_addr or '127.0.0.1').split(',')[0].strip()
            except Exception:
                ip = '127.0.0.1'

        uname = username or user_info.get('username') or 'Anonymous'
        urole = role or user_info.get('role') or 'N/A'

        entity = {
            'PartitionKey': 'log',
            'RowKey': row_key,
            'timestamp': now.isoformat(),
            'timestamp_formatted': now.strftime("%Y-%m-%d %H:%M:%S UTC"),
            'username': uname,
            'role': urole,
            'service': service,
            'action': action,
            'target': str(target)[:500] if target else '',
            'status': status,
            'details': str(details)[:1000] if details else '',
            'ip_address': ip
        }
        table_client.create_entity(entity)
    except Exception as e:
        # Never crash application if logging fails
        print(f"Activity logging error: {e}")

def cleanup_old_activity_logs(retention_days=None):
    """Purges activity logs older than retention_days (default ACTIVITY_LOG_RETENTION_DAYS=30 days) from LOGS_TABLE."""
    try:
        if not is_storage_connected():
            return 0
        days = retention_days if retention_days is not None else ACTIVITY_LOG_RETENTION_DAYS
        table_svc = get_table_service()
        table_client = table_svc.get_table_client(LOGS_TABLE)

        cutoff_dt = datetime.now(timezone.utc) - timedelta(days=days)
        cutoff_iso = cutoff_dt.isoformat()

        try:
            entities = list(table_client.query_entities("PartitionKey eq 'log'"))
        except Exception:
            return 0

        deleted_count = 0
        for ent in entities:
            ts_str = ent.get('timestamp')
            if ts_str:
                try:
                    log_dt = datetime.fromisoformat(ts_str.replace('Z', '+00:00'))
                    if log_dt < cutoff_dt:
                        table_client.delete_entity(partition_key='log', row_key=ent['RowKey'])
                        deleted_count += 1
                except Exception:
                    pass
        return deleted_count
    except Exception as e:
        print(f"Error cleaning up old activity logs: {e}")
        return 0

def query_activity_logs(service=None, username=None, status=None, from_date=None, to_date=None, limit=250):
    try:
        table_svc = get_table_service()
        table_client = table_svc.get_table_client(LOGS_TABLE)
        entities = table_client.query_entities("PartitionKey eq 'log'", results_per_page=limit * 2)
        logs = []

        uname_filter = (username or '').lower().strip()
        svc_filter = (service or 'all').lower().strip()
        stat_filter = (status or 'all').strip()

        # Enforce maximum 30 days retention policy
        retention_cutoff_dt = (datetime.now(timezone.utc) - timedelta(days=ACTIVITY_LOG_RETENTION_DAYS)).date()

        from_dt = None
        to_dt = None
        if from_date:
            try:
                parsed_from = datetime.strptime(from_date, "%Y-%m-%d").date()
                from_dt = max(parsed_from, retention_cutoff_dt)
            except Exception:
                from_dt = retention_cutoff_dt
        else:
            from_dt = retention_cutoff_dt

        if to_date:
            try:
                to_dt = datetime.strptime(to_date, "%Y-%m-%d").date()
            except Exception:
                pass

        for e in entities:
            l = dict(e)
            if svc_filter and svc_filter != 'all' and l.get('service', '').lower() != svc_filter:
                continue
            if uname_filter and uname_filter not in l.get('username', '').lower():
                continue
            if stat_filter and stat_filter != 'all' and l.get('status', '') != stat_filter:
                continue

            ts_str = l.get('timestamp', '')
            if ts_str:
                try:
                    log_date = datetime.fromisoformat(ts_str.replace('Z', '+00:00')).date()
                    if from_dt and log_date < from_dt:
                        continue
                    if to_dt and log_date > to_dt:
                        continue
                except Exception:
                    pass
            logs.append(l)
            if len(logs) >= limit:
                break
        return logs
    except Exception as e:
        print(f"Error querying activity logs: {e}")
        return []

# -----------------------
# Sidebar Navigation Tree Helper
# -----------------------
def load_sidebar_tree():
    if not require_auth():
        return {}
    
    tree = {
        'containers': [],
        'shares': [],
        'queues': [],
        'tables': []
    }
    
    # 1. Containers
    try:
        blob_svc = get_blob_service()
        tree['containers'] = [c.name for c in blob_svc.list_containers(results_per_page=100)]
    except Exception as e:
        print(f"Error loading containers for sidebar: {e}")
        
    # 2. File Shares
    try:
        share_svc = get_share_service()
        tree['shares'] = [s.name for s in share_svc.list_shares(results_per_page=100)]
    except Exception as e:
        print(f"Error loading shares for sidebar: {e}")
        
    # 3. Queues
    try:
        queue_svc = get_queue_service()
        tree['queues'] = [q.name for q in queue_svc.list_queues(results_per_page=100)]
    except Exception as e:
        print(f"Error loading queues for sidebar: {e}")
        
    # 4. Tables (Hide system tables from non-admin users)
    try:
        table_svc = get_table_service()
        tables = list(table_svc.list_tables(results_per_page=100))
        table_names = [t.name if hasattr(t, "name") else str(t) for t in tables]
        if not is_admin():
            table_names = [t for t in table_names if t.lower() not in [USER_TABLE.lower(), LOGS_TABLE.lower()]]
        tree['tables'] = table_names
    except Exception as e:
        print(f"Error loading tables for sidebar: {e}")
        
    return tree

# -----------------------
# Authentication & Portal Setup Routes
# -----------------------
@ui.route('/', methods=['GET', 'POST'])
def login():
    auto_connect_from_env_if_available()
    storage_connected = is_storage_connected()
    is_env_managed = bool(os.environ.get('AZURE_STORAGE_CONNECTION_STRING') or os.environ.get('AZURE_STORAGE_ACCOUNT_NAME'))

    # If storage connected, check if any users exist
    if storage_connected:
        ensure_system_tables_exist()
        if not check_has_any_users():
            return redirect(url_for('ui.setup'))

    if request.method == 'POST':
        login_type = request.form.get('login_type')
        
        # User Authentication (Username + Password)
        if login_type == 'user_auth' or ('username' in request.form and 'password' in request.form and not request.form.get('conn_string')):
            if not storage_connected:
                flash("Storage connection is not configured. Please connect to Azure Storage first.")
                return redirect(url_for('ui.login'))
                
            username = request.form.get('username', '').strip()
            password = request.form.get('password', '')
            
            user = get_user_by_username(username)
            if not user or not check_password_hash(user.get('password_hash', ''), password):
                log_activity('auth', 'LOGIN_FAILED', username, status='FAILED', details='Invalid username or password', username=username, role='N/A')
                flash("Invalid username or password.")
                return redirect(url_for('ui.login'))
                
            if not user.get('is_active', True):
                log_activity('auth', 'LOGIN_BLOCKED', username, status='FAILED', details='Account is disabled', username=username, role=user.get('role'))
                flash("Your account is disabled. Please contact your administrator.")
                return redirect(url_for('ui.login'))
                
            # Check if password reset is enforced
            if user.get('must_change_password', False):
                session['pending_user'] = username
                return redirect(url_for('ui.force_password_reset'))

            # Check password expiry policy
            try:
                expiry_days = int(user.get('password_expiry_days', DEFAULT_PASSWORD_EXPIRY_DAYS) or 0)
            except Exception:
                expiry_days = DEFAULT_PASSWORD_EXPIRY_DAYS

            if expiry_days > 0:
                last_set = user.get('password_last_set_at') or user.get('created_at')
                if last_set:
                    try:
                        last_set_dt = datetime.fromisoformat(last_set.replace('Z', '+00:00'))
                        age_days = (datetime.now(timezone.utc) - last_set_dt).days
                        if age_days >= expiry_days:
                            save_user(username=username, must_change_password=True)
                            session['pending_user'] = username
                            flash(f"Your password has expired ({age_days} days old, max policy {expiry_days} days). Please create a new password.")
                            return redirect(url_for('ui.force_password_reset'))
                    except Exception as e:
                        print(f"Password expiry check error: {e}")
                
            # Update last login & establish authenticated session
            save_user(username=username, update_login=True)
            session['user'] = {
                'username': user['username'],
                'display_name': user.get('display_name') or user['username'],
                'email': user.get('email', ''),
                'role': user.get('role', 'contributor'),
                'is_authenticated': True
            }
            log_activity('auth', 'LOGIN', username, status='SUCCESS', username=username, role=user.get('role'))
            flash(f"Welcome back, {session['user']['display_name']}!")
            return redirect(url_for('ui.home'))
            
        else:
            # Storage backend connection submitted
            return connect_storage()

    if require_auth():
        return redirect(url_for('ui.home'))
        
    return render_template('login.html', storage_connected=storage_connected, is_env_managed=is_env_managed)

@ui.route('/connect-storage', methods=['POST'])
def connect_storage():
    auth_method = request.form.get('auth_method')
    try:
        if auth_method == 'conn_str':
            conn_str = request.form.get('conn_string', '').strip()
            if not conn_str:
                raise ValueError("Connection string required")
            client = BlobServiceClient.from_connection_string(conn_str)
            _ = list(client.list_containers(results_per_page=1))
            
            account_name = parse_account_name(conn_str)
            session['auth_method'] = 'Connection String'
            session['account_name'] = account_name
            session['conn_string'] = conn_str
            
        elif auth_method == 'key':
            account_name = request.form.get('account_name', '').strip()
            account_key = request.form.get('account_key', '').strip()
            if not account_name or not account_key:
                raise ValueError("Account name and key required")
            
            conn_str = f"DefaultEndpointsProtocol=https;AccountName={account_name};AccountKey={account_key};EndpointSuffix=core.windows.net"
            client = BlobServiceClient.from_connection_string(conn_str)
            _ = list(client.list_containers(results_per_page=1))
            
            session['auth_method'] = 'Access Key'
            session['account_name'] = account_name
            session['account_key'] = account_key
            
        elif auth_method == 'sp':
            account_name = request.form.get('account_name_sp', '').strip()
            tenant_id = request.form.get('tenant_id', '').strip()
            client_id = request.form.get('client_id', '').strip()
            client_secret = request.form.get('client_secret', '').strip()
            if not all([account_name, tenant_id, client_id, client_secret]):
                raise ValueError("All Service Principal fields are required")
            
            credential = ClientSecretCredential(tenant_id, client_id, client_secret)
            client = BlobServiceClient(account_url=f"https://{account_name}.blob.core.windows.net", credential=credential)
            _ = list(client.list_containers(results_per_page=1))
            
            session['auth_method'] = 'Service Principal'
            session['account_name'] = account_name
            session['tenant_id'] = tenant_id
            session['client_id'] = client_id
            session['client_secret'] = client_secret
        else:
            raise ValueError("Invalid authentication method")
            
        ensure_system_tables_exist()
        if not check_has_any_users():
            return redirect(url_for('ui.setup'))
            
        flash("Connected to Azure Storage. Please sign in.")
        return redirect(url_for('ui.login'))
        
    except Exception as e:
        flash(f"Connection failed: {e}")
        return redirect(url_for('ui.login'))

@ui.route('/setup', methods=['GET', 'POST'])
def setup():
    auto_connect_from_env_if_available()
    if not is_storage_connected():
        flash("Please configure storage backend first.")
        return redirect(url_for('ui.login'))
        
    ensure_system_tables_exist()
    if check_has_any_users():
        flash("Setup has already been completed.")
        return redirect(url_for('ui.login'))
        
    if request.method == 'POST':
        username = request.form.get('username', '').strip()
        email = request.form.get('email', '').strip()
        display_name = request.form.get('display_name', '').strip() or 'System Administrator'
        password = request.form.get('password', '')
        confirm_password = request.form.get('confirm_password', '')
        
        if not username or not password:
            flash("Username and password are required.")
            return render_template('setup.html')
        if password != confirm_password:
            flash("Passwords do not match.")
            return render_template('setup.html')
        if len(password) < 6:
            flash("Password must be at least 6 characters.")
            return render_template('setup.html')
            
        try:
            save_user(
                username=username,
                password=password,
                email=email,
                display_name=display_name,
                role='admin',
                is_active=True,
                must_change_password=False,
                update_login=True
            )
            session['user'] = {
                'username': username,
                'display_name': display_name,
                'email': email,
                'role': 'admin',
                'is_authenticated': True
            }
            log_activity('auth', 'INITIAL_SETUP', username, status='SUCCESS', details='Master Admin Created', username=username, role='admin')
            flash("Master administrator account initialized successfully!")
            return redirect(url_for('ui.home'))
        except Exception as e:
            flash(f"Failed to create admin account: {e}")
            return render_template('setup.html')
            
    return render_template('setup.html')

@ui.route('/force-password-reset', methods=['GET', 'POST'])
def force_password_reset():
    username = session.get('pending_user')
    if not username:
        return redirect(url_for('ui.login'))
        
    if request.method == 'POST':
        new_password = request.form.get('new_password', '')
        confirm_password = request.form.get('confirm_password', '')
        
        if len(new_password) < 6:
            flash("New password must be at least 6 characters.")
            return render_template('force_password_reset.html')
        if new_password != confirm_password:
            flash("Passwords do not match.")
            return render_template('force_password_reset.html')
            
        try:
            save_user(username=username, password=new_password, must_change_password=False, update_login=True)
            user = get_user_by_username(username)
            session.pop('pending_user', None)
            session['user'] = {
                'username': user['username'],
                'display_name': user.get('display_name') or user['username'],
                'email': user.get('email', ''),
                'role': user.get('role', 'contributor'),
                'is_authenticated': True
            }
            log_activity('auth', 'PASSWORD_RESET_FORCED', username, status='SUCCESS', username=username, role=user.get('role'))
            flash("Password updated successfully! Welcome to Storage UI.")
            return redirect(url_for('ui.home'))
        except Exception as e:
            flash(f"Error resetting password: {e}")
            return render_template('force_password_reset.html')
            
    return render_template('force_password_reset.html')

@ui.route('/change-my-password', methods=['POST'])
def change_my_password():
    if not require_auth():
        return redirect(url_for('ui.login'))
        
    username = session['user']['username']
    current_password = request.form.get('current_password', '')
    new_password = request.form.get('new_password', '')
    confirm_password = request.form.get('confirm_password', '')
    
    if len(new_password) < 4:
        flash("New password must be at least 4 characters.")
        return redirect(request.referrer or url_for('ui.home'))
    if new_password != confirm_password:
        flash("New passwords do not match.")
        return redirect(request.referrer or url_for('ui.home'))
        
    user = get_user_by_username(username)
    if not user or not check_password_hash(user.get('password_hash', ''), current_password):
        flash("Current password incorrect.")
        return redirect(request.referrer or url_for('ui.home'))
        
    try:
        save_user(username=username, password=new_password, must_change_password=False)
        log_activity('auth', 'PASSWORD_CHANGE', username, status='SUCCESS', username=username, role=session['user']['role'])
        flash("Your password has been changed successfully.")
    except Exception as e:
        flash(f"Error changing password: {e}")
        
    return redirect(request.referrer or url_for('ui.home'))

@ui.route('/logout')
def logout():
    uname = session.get('user', {}).get('username', 'Anonymous')
    log_activity('auth', 'LOGOUT', uname, status='SUCCESS')
    session.pop('user', None)
    flash("Signed out successfully.")
    return redirect(url_for('ui.login'))

@ui.route('/home')
def home():
    if not require_auth():
        return redirect(url_for('ui.login'))
    tree = load_sidebar_tree()
    return render_template('dashboard.html', sidebar_tree=tree, active_service=None, active_item=None)

# -----------------------
# Portal Management: User & Role Management Routes (Admin Only)
# -----------------------
@ui.route('/users')
def users_list():
    if not is_admin():
        flash("Permission denied. Administrator role required.")
        return redirect(url_for('ui.home'))
        
    users = get_all_users()
    tree = load_sidebar_tree()
    return render_template('users.html', users=users, user_table_name=USER_TABLE, sidebar_tree=tree, active_service='users', active_item=None)

@ui.route('/users/create', methods=['POST'])
def create_user():
    if not is_admin():
        flash("Permission denied. Administrator role required.")
        return redirect(url_for('ui.home'))
        
    username = request.form.get('username', '').strip()
    display_name = request.form.get('display_name', '').strip()
    email = request.form.get('email', '').strip()
    role = request.form.get('role', 'contributor').strip().lower()
    password = request.form.get('password', '')
    enforce_reset = bool(request.form.get('enforce_reset'))
    expiry_days_raw = request.form.get('password_expiry_days')
    try:
        expiry_days = int(expiry_days_raw) if expiry_days_raw is not None and expiry_days_raw != '' else DEFAULT_PASSWORD_EXPIRY_DAYS
    except Exception:
        expiry_days = DEFAULT_PASSWORD_EXPIRY_DAYS
    
    if not username or not password:
        flash("Username and password are required.")
        return redirect(url_for('ui.users_list'))
        
    existing = get_user_by_username(username)
    if existing:
        flash(f"User '{username}' already exists.")
        return redirect(url_for('ui.users_list'))
        
    try:
        save_user(
            username=username,
            password=password,
            email=email,
            display_name=display_name,
            role=role,
            is_active=True,
            must_change_password=enforce_reset,
            password_expiry_days=expiry_days
        )
        log_activity('user_mgmt', 'CREATE_USER', username, status='SUCCESS', details=f"Role: {role}, Email: {email}, Expiry: {expiry_days}d")
        flash(f"User '{username}' created successfully.")
    except Exception as e:
        log_activity('user_mgmt', 'CREATE_USER', username, status='FAILED', details=str(e))
        flash(f"Failed to create user: {e}")
        
    return redirect(url_for('ui.users_list'))

@ui.route('/users/bulk-preview', methods=['POST'])
def bulk_preview_users():
    if not is_admin():
        return jsonify({"success": False, "error": "Permission denied. Administrator role required."}), 403
        
    uploaded_file = request.files.get('file')
    if not uploaded_file or not uploaded_file.filename:
        return jsonify({"success": False, "error": "No file was uploaded."}), 400
        
    parsed_users, summary, errors = parse_users_file(uploaded_file, uploaded_file.filename)
    if not parsed_users and errors:
        return jsonify({"success": False, "error": "; ".join(errors)}), 400

    # Store raw parsed list in session for one-click confirmation
    session['bulk_preview_users'] = parsed_users
    
    # Return sanitized view payload with masked passwords
    preview_users = []
    for u in parsed_users:
        preview_users.append({
            'row_index': u['row_index'],
            'username': u['username'],
            'display_name': u['display_name'],
            'email': u['email'],
            'role': u['role'],
            'masked_password': u['masked_password'],
            'must_change_password': u['must_change_password'],
            'password_expiry_days': u['password_expiry_days'],
            'is_valid': u['is_valid'],
            'already_exists': u['already_exists'],
            'error_msg': u['error_msg']
        })

    return jsonify({
        "success": True,
        "summary": summary,
        "users": preview_users,
        "errors": errors
    })

@ui.route('/users/bulk-create', methods=['POST'])
def bulk_create_users():
    if not is_admin():
        if request.headers.get('X-Requested-With') == 'XMLHttpRequest' or request.is_json:
            return jsonify({"success": False, "error": "Permission denied."}), 403
        flash("Permission denied. Administrator role required.")
        return redirect(url_for('ui.home'))
        
    is_ajax = request.headers.get('X-Requested-With') == 'XMLHttpRequest' or request.is_json
    data = request.get_json(silent=True) or {}
    confirm = data.get('confirm') or request.form.get('confirm')

    if confirm:
        # Create users from previously parsed preview session cache
        cached_users = session.pop('bulk_preview_users', None)
        if not cached_users:
            if is_ajax:
                return jsonify({"success": False, "error": "Preview session expired. Please re-upload the file."}), 400
            flash("Preview session expired. Please re-upload the file.")
            return redirect(url_for('ui.users_list'))
            
        created = 0
        skipped = 0
        errors = []
        for u in cached_users:
            if not u['is_valid']:
                skipped += 1
                continue
            try:
                save_user(
                    username=u['username'],
                    password=u['password'],
                    email=u['email'],
                    display_name=u['display_name'],
                    role=u['role'],
                    is_active=True,
                    must_change_password=u['must_change_password'],
                    password_expiry_days=u.get('password_expiry_days', DEFAULT_PASSWORD_EXPIRY_DAYS)
                )
                created += 1
            except Exception as e:
                skipped += 1
                errors.append(f"Row {u['row_index']} ({u['username']}): {e}")

        details = f"Created: {created}, Skipped: {skipped}"
        if errors:
            details += f" (Errors: {'; '.join(errors[:3])})"
        log_activity('user_mgmt', 'BULK_CREATE_USERS', 'Bulk Import (Preview Confirmed)', status='SUCCESS' if created > 0 else 'FAILED', details=details)

        msg = f"Bulk import complete: {created} user(s) created."
        if skipped > 0:
            msg += f" {skipped} skipped."
        if is_ajax:
            return jsonify({"success": True, "created": created, "skipped": skipped, "errors": errors, "message": msg})
        flash(msg)
        return redirect(url_for('ui.users_list'))

    # Fallback direct file upload without preview
    uploaded_file = request.files.get('file')
    if not uploaded_file or not uploaded_file.filename:
        if is_ajax:
            return jsonify({"success": False, "error": "No file was uploaded."}), 400
        flash("No file was uploaded.")
        return redirect(url_for('ui.users_list'))
        
    created, skipped, errors = bulk_create_users_from_file(uploaded_file, uploaded_file.filename)
    details = f"Created: {created}, Skipped: {skipped}"
    if errors:
        details += f" (Errors: {'; '.join(errors[:3])})"
        
    log_activity('user_mgmt', 'BULK_CREATE_USERS', uploaded_file.filename, status='SUCCESS' if created > 0 else 'FAILED', details=details)
    
    msg = f"Bulk import complete: {created} user(s) created."
    if skipped > 0:
        msg += f" {skipped} skipped. {'; '.join(errors)}"
    if is_ajax:
        return jsonify({"success": True, "created": created, "skipped": skipped, "errors": errors, "message": msg})
    flash(msg)
    return redirect(url_for('ui.users_list'))

@ui.route('/users/edit', methods=['POST'])
def edit_user():
    if not is_admin():
        flash("Permission denied. Administrator role required.")
        return redirect(url_for('ui.home'))
        
    username = request.form.get('username', '').strip()
    display_name = request.form.get('display_name', '').strip()
    email = request.form.get('email', '').strip()
    role = request.form.get('role', 'contributor').strip().lower()
    expiry_days_raw = request.form.get('password_expiry_days')
    try:
        expiry_days = int(expiry_days_raw) if expiry_days_raw is not None and expiry_days_raw != '' else None
    except Exception:
        expiry_days = None
    
    try:
        save_user(username=username, display_name=display_name, email=email, role=role, password_expiry_days=expiry_days)
        # If editing self, update current session
        if session.get('user', {}).get('username') == username:
            session['user']['display_name'] = display_name or username
            session['user']['email'] = email
            session['user']['role'] = role
            
        log_activity('user_mgmt', 'EDIT_USER', username, status='SUCCESS', details=f"Role: {role}, Email: {email}, Expiry: {expiry_days}d")
        flash(f"User '{username}' updated successfully.")
    except Exception as e:
        log_activity('user_mgmt', 'EDIT_USER', username, status='FAILED', details=str(e))
        flash(f"Failed to update user: {e}")
        
    return redirect(url_for('ui.users_list'))

@ui.route('/users/reset-password', methods=['POST'])
def reset_user_password():
    if not is_admin():
        flash("Permission denied. Administrator role required.")
        return redirect(url_for('ui.home'))
        
    username = request.form.get('username', '').strip()
    new_password = request.form.get('new_password', '')
    enforce_reset = bool(request.form.get('enforce_reset'))
    
    if len(new_password) < 4:
        flash("Password must be at least 4 characters.")
        return redirect(url_for('ui.users_list'))
        
    try:
        save_user(username=username, password=new_password, must_change_password=enforce_reset)
        log_activity('user_mgmt', 'RESET_PASSWORD', username, status='SUCCESS', details=f"Enforce Reset: {enforce_reset}")
        flash(f"Password for '{username}' has been reset.")
    except Exception as e:
        log_activity('user_mgmt', 'RESET_PASSWORD', username, status='FAILED', details=str(e))
        flash(f"Failed to reset password: {e}")
        
    return redirect(url_for('ui.users_list'))

@ui.route('/users/toggle-status', methods=['POST'])
def toggle_user_status():
    if not is_admin():
        flash("Permission denied. Administrator role required.")
        return redirect(url_for('ui.home'))
        
    username = request.form.get('username', '').strip()
    is_active_val = request.form.get('is_active', '1') == '1'
    
    if username == session.get('user', {}).get('username') and not is_active_val:
        flash("You cannot disable your own administrator account.")
        return redirect(url_for('ui.users_list'))
        
    try:
        save_user(username=username, is_active=is_active_val)
        status_name = "enabled" if is_active_val else "disabled"
        log_activity('user_mgmt', f"{'ENABLE' if is_active_val else 'DISABLE'}_USER", username, status='SUCCESS')
        flash(f"Account for '{username}' has been {status_name}.")
    except Exception as e:
        log_activity('user_mgmt', 'TOGGLE_STATUS_USER', username, status='FAILED', details=str(e))
        flash(f"Failed to update status: {e}")
        
    return redirect(url_for('ui.users_list'))

@ui.route('/users/delete', methods=['POST'])
def delete_user():
    if not is_admin():
        flash("Permission denied. Administrator role required.")
        return redirect(url_for('ui.home'))
        
    username = request.form.get('username', '').strip()
    if username == session.get('user', {}).get('username'):
        flash("You cannot delete your own administrator account.")
        return redirect(url_for('ui.users_list'))
        
    try:
        delete_user_by_username(username)
        log_activity('user_mgmt', 'DELETE_USER', username, status='SUCCESS')
        flash(f"User '{username}' deleted successfully.")
    except Exception as e:
        log_activity('user_mgmt', 'DELETE_USER', username, status='FAILED', details=str(e))
        flash(f"Failed to delete user: {e}")
        
    return redirect(url_for('ui.users_list'))

@ui.route('/users/template.csv')
def download_user_template():
    csv_data = "username,email,display_name,password,enforcepasswordreset,role,password_expiry_days\n"
    csv_data += "john,john@example.com,John Doe,Password123!,yes,contributor,90\n"
    csv_data += "sarah,sarah@example.com,Sarah Smith,TempPass456!,yes,reader,90\n"
    csv_data += "admin2,admin2@example.com,Secondary Admin,SecureAdmin789!,no,admin,0\n"
    return Response(
        csv_data,
        mimetype="text/csv",
        headers={"Content-Disposition": "attachment;filename=users_import_template.csv"}
    )

# -----------------------
# Portal Management: Activity Logs Routes (Admin Only)
# -----------------------
@ui.route('/activity-logs')
def activity_logs():
    if not is_admin():
        flash("Permission denied. Administrator role required.")
        return redirect(url_for('ui.home'))
        
    # Auto-prune logs older than 30 days
    cleanup_old_activity_logs()
    
    filter_service = request.args.get('service', 'all')
    filter_username = request.args.get('username', '')
    filter_status = request.args.get('status', 'all')
    from_date = request.args.get('from_date', '')
    to_date = request.args.get('to_date', '')
    
    logs = query_activity_logs(
        service=filter_service,
        username=filter_username,
        status=filter_status,
        from_date=from_date,
        to_date=to_date,
        limit=250
    )
    tree = load_sidebar_tree()
    return render_template(
        'activity_logs.html',
        logs=logs,
        filter_service=filter_service,
        filter_username=filter_username,
        filter_status=filter_status,
        from_date=from_date,
        to_date=to_date,
        sidebar_tree=tree,
        active_service='activity_logs',
        active_item=None
    )

@ui.route('/activity-logs/export')
def export_activity_logs():
    if not is_admin():
        abort(403)
        
    fmt = request.args.get('format', 'csv').lower()
    filter_service = request.args.get('service', 'all')
    filter_username = request.args.get('username', '')
    filter_status = request.args.get('status', 'all')
    from_date = request.args.get('from_date', '')
    to_date = request.args.get('to_date', '')
    
    logs = query_activity_logs(
        service=filter_service,
        username=filter_username,
        status=filter_status,
        from_date=from_date,
        to_date=to_date,
        limit=1000
    )
    
    if fmt == 'json':
        clean_logs = []
        for l in logs:
            c = dict(l)
            c.pop('PartitionKey', None)
            c.pop('RowKey', None)
            c.pop('odata.etag', None)
            c.pop('etag', None)
            c.pop('Timestamp', None)
            clean_logs.append(c)
        return Response(
            json.dumps(clean_logs, indent=2),
            mimetype="application/json",
            headers={"Content-Disposition": "attachment;filename=activity_logs.json"}
        )
    else:
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow(["Timestamp", "Username", "Role", "Service", "Action", "Target", "Status", "Client IP", "Details"])
        for l in logs:
            writer.writerow([
                l.get('timestamp_formatted') or l.get('timestamp', ''),
                l.get('username', ''),
                l.get('role', ''),
                l.get('service', ''),
                l.get('action', ''),
                l.get('target', ''),
                l.get('status', ''),
                l.get('ip_address', ''),
                l.get('details', '')
            ])
        output.seek(0)
        return Response(
            output.getvalue(),
            mimetype="text/csv",
            headers={"Content-Disposition": "attachment;filename=activity_logs.csv"}
        )

# -----------------------
# Bulk Create Resource (Containers, Shares, Queues, Tables)
# -----------------------
@ui.route('/bulk-create', methods=['POST'])
def bulk_create():
    if not require_auth():
        return redirect(url_for('ui.login'))
    if not is_admin():
        flash("Permission denied. Administrator role required for resource creation.")
        return redirect(url_for('ui.home'))
        
    resource_type = request.form.get('resource_type')
    uploaded_file = request.files.get('file')
    
    if not uploaded_file or not uploaded_file.filename:
        flash("No file was uploaded.")
        return redirect(url_for('ui.home'))
        
    filename = uploaded_file.filename.lower()
    names = []
    
    try:
        if filename.endswith('.csv'):
            stream = io.StringIO(uploaded_file.stream.read().decode("utf-8-sig"))
            reader = csv.reader(stream)
            for row in reader:
                if row:
                    name = row[0].strip()
                    if name:
                        names.append(name)
        elif filename.endswith('.xlsx'):
            wb = openpyxl.load_workbook(io.BytesIO(uploaded_file.read()), data_only=True)
            sheet = wb.active
            for row in sheet.iter_rows(min_row=1, max_col=1, values_only=True):
                if row and row[0] is not None:
                    name = str(row[0]).strip()
                    if name:
                        names.append(name)
        else:
            flash("Unsupported file format. Please upload a .csv or .xlsx file.")
            return redirect(url_for('ui.home'))
    except Exception as e:
        flash(f"Error parsing file: {e}")
        return redirect(url_for('ui.home'))
        
    header_keywords = {"name", "names", "container", "containers", "share", "shares", "table", "tables", "queue", "queues"}
    filtered_names = []
    for n in names:
        if n.lower() not in header_keywords:
            filtered_names.append(n)
            
    if not filtered_names:
        flash("No valid resource names found in the uploaded file.")
        return redirect(url_for('ui.home'))
        
    success_count = 0
    errors = []
    
    try:
        if resource_type == 'container':
            svc = get_blob_service()
            for name in filtered_names:
                try:
                    svc.create_container(name)
                    log_activity('blob', 'CREATE_CONTAINER', name, 'SUCCESS')
                    success_count += 1
                except Exception as ex:
                    errors.append(f"'{name}': {ex}")
                    log_activity('blob', 'CREATE_CONTAINER', name, 'FAILED', details=str(ex))
                    
        elif resource_type == 'share':
            svc = get_share_service()
            for name in filtered_names:
                try:
                    svc.create_share(name)
                    log_activity('file', 'CREATE_SHARE', name, 'SUCCESS')
                    success_count += 1
                except Exception as ex:
                    errors.append(f"'{name}': {ex}")
                    log_activity('file', 'CREATE_SHARE', name, 'FAILED', details=str(ex))
                    
        elif resource_type == 'queue':
            svc = get_queue_service()
            for name in filtered_names:
                try:
                    svc.create_queue(name)
                    log_activity('queue', 'CREATE_QUEUE', name, 'SUCCESS')
                    success_count += 1
                except Exception as ex:
                    errors.append(f"'{name}': {ex}")
                    log_activity('queue', 'CREATE_QUEUE', name, 'FAILED', details=str(ex))
                    
        elif resource_type == 'table':
            svc = get_table_service()
            for name in filtered_names:
                try:
                    svc.create_table(name)
                    log_activity('table', 'CREATE_TABLE', name, 'SUCCESS')
                    success_count += 1
                except Exception as ex:
                    errors.append(f"'{name}': {ex}")
                    log_activity('table', 'CREATE_TABLE', name, 'FAILED', details=str(ex))
        else:
            flash(f"Invalid resource type: {resource_type}")
            return redirect(url_for('ui.home'))
            
    except Exception as e:
        flash(f"Service connection error: {e}")
        return redirect(url_for('ui.home'))
        
    resource_type_display = {
        'container': 'Blob Container(s)',
        'share': 'File Share(s)',
        'queue': 'Message Queue(s)',
        'table': 'NoSQL Table(s)'
    }.get(resource_type, 'Resource(s)')
    
    summary = f"Bulk creation completed. Successfully created {success_count} of {len(filtered_names)} {resource_type_display}."
    if errors:
        summary += f" Errors: {'; '.join(errors)}"
    flash(summary)
    
    return redirect(url_for('ui.home'))

# -----------------------
# Blob Routes
# -----------------------
@ui.route('/blobs')
def blobs():
    if not require_auth():
        return redirect(url_for('ui.login'))

    try:
        blob_service = get_blob_service()
        containers = list(blob_service.list_containers())
    except Exception as e:
        flash(f"Error loading containers: {e}")
        containers = []
        
    tree = load_sidebar_tree()
    return render_template('blobs.html', containers=containers, container_name=None, sidebar_tree=tree, active_service='blobs', active_item=None)

@ui.route('/blobs/create', methods=['POST'])
def create_container():
    if not require_auth():
        return redirect(url_for('ui.login'))
    if not is_admin():
        flash("Permission denied. Administrator role required to create containers.")
        return redirect(url_for('ui.blobs'))
        
    name = request.form.get('name') or request.form.get('container_name')
    try:
        get_blob_service().create_container(name)
        log_activity('blob', 'CREATE_CONTAINER', name, 'SUCCESS')
        flash(f"Container '{name}' created.")
    except Exception as e:
        log_activity('blob', 'CREATE_CONTAINER', name, 'FAILED', details=str(e))
        flash(f"Error creating container: {e}")
    return redirect(url_for('ui.blobs'))

@ui.route('/blobs/delete', methods=['POST'])
def delete_container():
    if not require_auth():
        return redirect(url_for('ui.login'))
    if not is_admin():
        flash("Permission denied. Administrator role required to delete containers.")
        return redirect(url_for('ui.blobs'))
        
    name = request.form.get('container_name')
    try:
        get_blob_service().delete_container(name)
        log_activity('blob', 'DELETE_CONTAINER', name, 'SUCCESS')
        flash(f"Container '{name}' deleted.")
    except Exception as e:
        log_activity('blob', 'DELETE_CONTAINER', name, 'FAILED', details=str(e))
        flash(f"Error deleting container: {e}")
    return redirect(url_for('ui.blobs'))

@ui.route('/blobs/<container_name>', methods=['GET', 'POST'])
def view_blobs(container_name):
    if not require_auth():
        if request.headers.get('X-Requested-With') == 'XMLHttpRequest' or request.is_json:
            return jsonify({"success": False, "error": "Not authenticated"}), 401
        return redirect(url_for('ui.login'))

    service = get_blob_service()
    container_client = service.get_container_client(container_name)

    # Handle multiple files upload
    if request.method == 'POST' and 'files' in request.files:
        if not has_write_permission():
            if request.headers.get('X-Requested-With') == 'XMLHttpRequest' or request.is_json:
                return jsonify({"success": False, "error": "Permission denied. Reader role is read-only."}), 403
            flash("Permission denied. Reader role is read-only.")
            return redirect(url_for('ui.view_blobs', container_name=container_name))
            
        folder = request.form.get('folder', '').strip()
        files = request.files.getlist('files')

        uploaded = []
        failed = []
        for file in files:
            if not file or not file.filename:
                continue
            
            blob_name = file.filename
            if folder:
                blob_name = f"{folder.rstrip('/')}/{file.filename}"

            try:
                guessed_type = mimetypes.guess_type(blob_name)[0] or 'application/octet-stream'
                container_client.upload_blob(
                    name=blob_name,
                    data=file,
                    overwrite=True,
                    content_settings=ContentSettings(content_type=guessed_type)
                )
                uploaded.append(file.filename)
                log_activity('blob', 'UPLOAD_BLOB', f"{container_name}/{blob_name}", 'SUCCESS')
            except Exception as e:
                failed.append(f"{file.filename} ({e})")
                log_activity('blob', 'UPLOAD_BLOB', f"{container_name}/{blob_name}", 'FAILED', details=str(e))
                
        is_ajax = request.headers.get('X-Requested-With') == 'XMLHttpRequest' or request.headers.get('Accept') == 'application/json'
        if is_ajax:
            return jsonify({
                "success": len(uploaded) > 0 or len(failed) == 0,
                "uploaded_count": len(uploaded),
                "uploaded": uploaded,
                "failed": failed
            })

        if uploaded:
            flash(f"Successfully uploaded {len(uploaded)} file(s).")
        if failed:
            flash(f"Failed to upload: {', '.join(failed)}")

        return redirect(url_for('ui.view_blobs', container_name=container_name))

    # Pagination & Search & Sort parameters
    search_query = request.args.get('q', '').strip()
    try:
        page = max(1, int(request.args.get('page', 1)))
    except ValueError:
        page = 1
    try:
        limit = int(request.args.get('limit', 20))
        if limit not in [10, 15, 20, 50, 100]:
            limit = 20
    except ValueError:
        limit = 20

    sort_by = request.args.get('sort', 'name').strip().lower()
    order = request.args.get('order', 'asc').strip().lower()
    from_date = request.args.get('from_date', '').strip()
    to_date = request.args.get('to_date', '').strip()

    from_dt = None
    to_dt = None
    if from_date:
        try:
            from_dt = datetime.strptime(from_date, "%Y-%m-%d").date()
        except Exception:
            pass
    if to_date:
        try:
            to_dt = datetime.strptime(to_date, "%Y-%m-%d").date()
        except Exception:
            pass

    try:
        raw_blobs = list(container_client.list_blobs())
    except Exception as e:
        flash(f"Error fetching blobs: {e}")
        raw_blobs = []

    filtered_blobs = []
    for b in raw_blobs:
        if search_query and search_query.lower() not in b.name.lower():
            continue
        if from_dt or to_dt:
            if hasattr(b, 'last_modified') and b.last_modified:
                blob_date = b.last_modified.date()
                if from_dt and blob_date < from_dt:
                    continue
                if to_dt and blob_date > to_dt:
                    continue
        filtered_blobs.append(b)

    # Sorting
    def sort_key(b):
        if sort_by == 'size':
            return b.size or 0
        elif sort_by == 'date':
            return b.last_modified.timestamp() if getattr(b, 'last_modified', None) else 0
        elif sort_by == 'type':
            ct = getattr(b, 'content_settings', None)
            return (ct.content_type if ct and ct.content_type else '') or ''
        return b.name.lower()

    reverse = (order == 'desc')
    filtered_blobs.sort(key=sort_key, reverse=reverse)

    total_items = len(filtered_blobs)
    total_pages = max(1, math.ceil(total_items / limit))
    if page > total_pages:
        page = total_pages

    start_idx = (page - 1) * limit
    end_idx = start_idx + limit
    paginated_blobs = filtered_blobs[start_idx:end_idx]

    tree = load_sidebar_tree()
    return render_template(
        'blobs.html',
        blobs=paginated_blobs,
        container_name=container_name,
        containers=None,
        search_query=search_query,
        page=page,
        total_pages=total_pages,
        limit=limit,
        total_items=total_items,
        sort_by=sort_by,
        order=order,
        from_date=from_date,
        to_date=to_date,
        sidebar_tree=tree,
        active_service='blobs',
        active_item=container_name
    )

@ui.route('/blobs/<container_name>/delete-multiple', methods=['POST'])
def delete_multiple_blobs(container_name):
    if not require_auth():
        return jsonify({"success": False, "error": "Not authenticated"}), 401
    if not has_write_permission():
        return jsonify({"success": False, "error": "Permission denied. Reader role is read-only."}), 403

    data = request.get_json(silent=True) or request.form
    blob_names = data.get('blobs', [])
    if isinstance(blob_names, str):
        blob_names = [b.strip() for b in blob_names.split(',') if b.strip()]

    if not blob_names:
        return jsonify({"success": False, "error": "No blobs specified for deletion"}), 400

    service = get_blob_service()
    container_client = service.get_container_client(container_name)

    deleted = []
    failed = []
    for name in blob_names:
        try:
            container_client.delete_blob(name)
            deleted.append(name)
        except Exception as e:
            failed.append(f"{name}: {e}")

    status = 'SUCCESS' if len(deleted) > 0 else 'FAILED'
    log_activity('blob', 'DELETE_MULTIPLE_BLOBS', container_name, status, details=f"Deleted {len(deleted)} of {len(blob_names)}: {', '.join(deleted[:5])}")

    return jsonify({
        "success": len(deleted) > 0,
        "deleted_count": len(deleted),
        "failed_count": len(failed),
        "deleted": deleted,
        "failed": failed
    })

@ui.route('/blobs/<container_name>/content', methods=['GET'])
def get_blob_content(container_name):
    if not require_auth():
        return jsonify({"success": False, "error": "Not authenticated"}), 401

    blob_name = (request.args.get('blob_name') or request.args.get('name') or request.args.get('blob') or '').strip()
    if not blob_name:
        return jsonify({"success": False, "error": "Blob name required"}), 400

    service = get_blob_service()
    container_client = service.get_container_client(container_name)
    blob_client = container_client.get_blob_client(blob_name)

    try:
        props = blob_client.get_blob_properties()
        content_type = props.content_settings.content_type or mimetypes.guess_type(blob_name)[0] or 'application/octet-stream'
        
        # Max preview size: 10MB
        if props.size > 10 * 1024 * 1024:
            return jsonify({
                "success": False,
                "error": f"File is too large for inline preview ({props.size / (1024*1024):.1f} MB). Please download it instead."
            }), 400

        stream = blob_client.download_blob()
        raw_bytes = stream.readall()

        is_text = False
        text_content = ""
        is_image = content_type.startswith('image/')
        is_pdf = content_type == 'application/pdf'

        if not is_image and not is_pdf:
            try:
                text_content = raw_bytes.decode('utf-8')
                is_text = True
            except UnicodeDecodeError:
                is_text = False

        b64_data = ""
        if is_image or is_pdf or not is_text:
            b64_data = base64.b64encode(raw_bytes).decode('utf-8')

        return jsonify({
            "success": True,
            "name": blob_name,
            "size": props.size,
            "content_type": content_type,
            "is_text": is_text,
            "is_image": is_image,
            "is_pdf": is_pdf,
            "content": text_content if is_text else b64_data,
            "last_modified": props.last_modified.isoformat() if props.last_modified else None
        })
    except Exception as e:
        return jsonify({"success": False, "error": str(e)}), 400

@ui.route('/blobs/<container_name>/save-content', methods=['POST'])
def save_blob_content(container_name):
    if not require_auth():
        return jsonify({"success": False, "error": "Not authenticated"}), 401
    if not has_write_permission():
        return jsonify({"success": False, "error": "Permission denied. Reader role is read-only."}), 403

    data = request.get_json(silent=True) or request.form
    blob_name = (data.get('blob_name') or data.get('name') or data.get('blob') or '').strip().lstrip('/')
    content = data.get('content')
    is_new = bool(data.get('is_new', False))

    if not blob_name or content is None:
        return jsonify({"success": False, "error": "Blob filename and content are required"}), 400

    try:
        service = get_blob_service()
        container_client = service.get_container_client(container_name)
        guessed_type = mimetypes.guess_type(blob_name)[0] or 'text/plain'
        container_client.upload_blob(
            name=blob_name,
            data=content.encode('utf-8'),
            overwrite=True,
            content_settings=ContentSettings(content_type=guessed_type)
        )
        action_name = 'CREATE_BLOB' if is_new else 'EDIT_BLOB_CONTENT'
        log_activity('blob', action_name, f"{container_name}/{blob_name}", 'SUCCESS')
        return jsonify({"success": True, "message": "Saved successfully", "name": blob_name})
    except Exception as e:
        action_name = 'CREATE_BLOB' if is_new else 'EDIT_BLOB_CONTENT'
        log_activity('blob', action_name, f"{container_name}/{blob_name}", 'FAILED', details=str(e))
        return jsonify({"success": False, "error": str(e)}), 400

@ui.route('/blobs/<container_name>/create-folder', methods=['POST'])
def create_blob_folder(container_name):
    if not require_auth():
        return jsonify({"success": False, "error": "Not authenticated"}), 401
    if not has_write_permission():
        return jsonify({"success": False, "error": "Permission denied. Reader role is read-only."}), 403

    data = request.get_json(silent=True) or request.form
    folder_path = data.get('folder_path', '').strip().strip('/')
    if not folder_path:
        return jsonify({"success": False, "error": "Folder path is required"}), 400

    placeholder = f"{folder_path}/.keep"
    try:
        service = get_blob_service()
        container_client = service.get_container_client(container_name)
        container_client.upload_blob(name=placeholder, data=b'', overwrite=True)
        log_activity('blob', 'CREATE_FOLDER', f"{container_name}/{folder_path}", 'SUCCESS')
        return jsonify({"success": True, "message": f"Folder '{folder_path}' created."})
    except Exception as e:
        log_activity('blob', 'CREATE_FOLDER', f"{container_name}/{folder_path}", 'FAILED', details=str(e))
        return jsonify({"success": False, "error": str(e)}), 400

@ui.route('/blobs/<container_name>/download')
def download_blob(container_name):
    if not require_auth():
        return redirect(url_for('ui.login'))
    name = request.args.get('name')
    service = get_blob_service()
    blob_client = service.get_container_client(container_name).get_blob_client(name)
    stream = io.BytesIO()
    blob_client.download_blob().readinto(stream)
    stream.seek(0)
    filename = secure_filename(name.split('/')[-1]) or 'download'
    return send_file(stream, download_name=filename, as_attachment=True)

@ui.route('/blobs/<container_name>/delete', methods=['POST'])
def delete_blob(container_name):
    if not require_auth():
        return redirect(url_for('ui.login'))
    if not has_write_permission():
        flash("Permission denied. Reader role is read-only.")
        return redirect(url_for('ui.view_blobs', container_name=container_name))
        
    name = request.form.get('blob_name')
    try:
        service = get_blob_service()
        service.get_container_client(container_name).delete_blob(name)
        log_activity('blob', 'DELETE_BLOB', f"{container_name}/{name}", 'SUCCESS')
        flash(f"Blob '{name}' deleted.")
    except Exception as e:
        log_activity('blob', 'DELETE_BLOB', f"{container_name}/{name}", 'FAILED', details=str(e))
        flash(f"Error deleting blob: {e}")
    return redirect(url_for('ui.view_blobs', container_name=container_name))

@ui.route('/blobs/<container_name>/download-selected', methods=['POST'])
def download_selected_blobs(container_name):
    if not require_auth():
        return jsonify({"success": False, "error": "Not authenticated"}), 401

    data = request.get_json(silent=True) or request.form
    blob_names = data.get('blobs', [])
    if isinstance(blob_names, str):
        blob_names = [b.strip() for b in blob_names.split(',') if b.strip()]

    if not blob_names:
        return jsonify({"success": False, "error": "No blobs selected"}), 400

    service = get_blob_service()
    container_client = service.get_container_client(container_name)

    # If single blob, download directly
    if len(blob_names) == 1:
        name = blob_names[0]
        try:
            blob_client = container_client.get_blob_client(name)
            stream = io.BytesIO()
            blob_client.download_blob().readinto(stream)
            stream.seek(0)
            filename = name.split('/')[-1] or 'download'
            return send_file(stream, download_name=filename, as_attachment=True)
        except Exception as e:
            return jsonify({"success": False, "error": str(e)}), 400

    # If multiple blobs, bundle into a ZIP archive
    zip_buffer = io.BytesIO()
    with zipfile.ZipFile(zip_buffer, 'w', zipfile.ZIP_DEFLATED) as zip_file:
        for name in blob_names:
            try:
                blob_client = container_client.get_blob_client(name)
                blob_data = blob_client.download_blob().readall()
                zip_path = name.lstrip('/')
                zip_file.writestr(zip_path, blob_data)
            except Exception as e:
                print(f"Failed to add {name} to zip: {e}")

    zip_buffer.seek(0)
    zip_filename = f"{secure_filename(container_name)}_selected_{datetime.now().strftime('%Y%m%d_%H%M%S')}.zip"
    return send_file(zip_buffer, mimetype='application/zip', download_name=zip_filename, as_attachment=True)

@ui.route('/blobs/<container_name>/download-all')
def download_all_blobs(container_name):
    if not require_auth():
        return redirect(url_for('ui.login'))

    service = get_blob_service()
    container_client = service.get_container_client(container_name)

    zip_buffer = io.BytesIO()
    count = 0
    with zipfile.ZipFile(zip_buffer, 'w', zipfile.ZIP_DEFLATED) as zip_file:
        for blob in container_client.list_blobs():
            try:
                blob_client = container_client.get_blob_client(blob.name)
                blob_data = blob_client.download_blob().readall()
                zip_path = blob.name.lstrip('/')
                zip_file.writestr(zip_path, blob_data)
                count += 1
            except Exception as e:
                print(f"Failed to add {blob.name} to zip: {e}")

    if count == 0:
        flash("No blobs available to download in this container.")
        return redirect(url_for('ui.view_blobs', container_name=container_name))

    zip_buffer.seek(0)
    zip_filename = f"{secure_filename(container_name)}_all_{datetime.now().strftime('%Y%m%d_%H%M%S')}.zip"
    return send_file(zip_buffer, mimetype='application/zip', download_name=zip_filename, as_attachment=True)

@ui.route('/blobs/<container_name>/empty', methods=['POST'])
def empty_container(container_name):
    if not require_auth():
        return redirect(url_for('ui.login'))
    if not has_write_permission():
        flash("Permission denied. Reader role is read-only.")
        return redirect(url_for('ui.view_blobs', container_name=container_name))

    service = get_blob_service()
    container_client = service.get_container_client(container_name)

    deleted_count = 0
    try:
        for blob in container_client.list_blobs():
            container_client.delete_blob(blob.name)
            deleted_count += 1
        log_activity('blob', 'EMPTY_CONTAINER', container_name, 'SUCCESS', details=f"Deleted {deleted_count} blobs")
        flash(f"Emptied container '{container_name}'. Deleted {deleted_count} blobs.")
    except Exception as e:
        log_activity('blob', 'EMPTY_CONTAINER', container_name, 'FAILED', details=str(e))
        flash(f"Error emptying container: {e}")

    return redirect(url_for('ui.view_blobs', container_name=container_name))

@ui.route('/blobs/<container_name>/rename', methods=['POST'])
def rename_blob(container_name):
    if not require_auth():
        return jsonify({"success": False, "error": "Not authenticated"}), 401
    if not has_write_permission():
        return jsonify({"success": False, "error": "Permission denied. Reader role is read-only."}), 403

    data = request.get_json(silent=True) or request.form
    old_name = data.get('old_name')
    new_name = data.get('new_name')

    if not old_name or not new_name:
        return jsonify({"success": False, "error": "Old and new blob names are required"}), 400

    service = get_blob_service()
    container_client = service.get_container_client(container_name)

    try:
        source_blob = container_client.get_blob_client(old_name)
        dest_blob = container_client.get_blob_client(new_name)

        dest_blob.start_copy_from_url(source_blob.url)
        source_blob.delete_blob()
        log_activity('blob', 'RENAME_BLOB', f"{container_name}/{old_name} -> {new_name}", 'SUCCESS')
        return jsonify({"success": True, "message": f"Renamed '{old_name}' to '{new_name}'."})
    except Exception as e:
        log_activity('blob', 'RENAME_BLOB', f"{container_name}/{old_name} -> {new_name}", 'FAILED', details=str(e))
        return jsonify({"success": False, "error": str(e)}), 400

@ui.route('/blobs/<container_name>/move-copy', methods=['POST'])
def move_copy_blob(container_name):
    if not require_auth():
        return jsonify({"success": False, "error": "Not authenticated"}), 401
    if not has_write_permission():
        return jsonify({"success": False, "error": "Permission denied. Reader role is read-only."}), 403

    data = request.get_json(silent=True) or request.form
    blob_name = data.get('blob_name')
    dest_container = data.get('dest_container')
    dest_name = data.get('dest_name') or blob_name
    action_type = data.get('action_type', 'copy')

    if not blob_name or not dest_container:
        return jsonify({"success": False, "error": "Blob name and destination container are required"}), 400

    service = get_blob_service()
    src_client = service.get_container_client(container_name).get_blob_client(blob_name)
    dest_client = service.get_container_client(dest_container).get_blob_client(dest_name)

    try:
        dest_client.start_copy_from_url(src_client.url)
        if action_type == 'move':
            src_client.delete_blob()
        log_activity('blob', f"{action_type.upper()}_BLOB", f"{container_name}/{blob_name} -> {dest_container}/{dest_name}", 'SUCCESS')
        return jsonify({"success": True, "message": f"Successfully {'moved' if action_type == 'move' else 'copied'} to '{dest_container}/{dest_name}'."})
    except Exception as e:
        log_activity('blob', f"{action_type.upper()}_BLOB", f"{container_name}/{blob_name} -> {dest_container}/{dest_name}", 'FAILED', details=str(e))
        return jsonify({"success": False, "error": str(e)}), 400

# -----------------------
# File Shares Routes
# -----------------------
@ui.route('/fileshares')
def fileshares():
    if not require_auth():
        return redirect(url_for('ui.login'))

    try:
        shares = list(get_share_service().list_shares())
    except Exception as e:
        flash(f"Error loading file shares: {e}")
        shares = []

    tree = load_sidebar_tree()
    return render_template('fileshares.html', shares=shares, share_name=None, sidebar_tree=tree, active_service='fileshares', active_item=None)

@ui.route('/fileshares/create', methods=['POST'])
def create_share():
    if not require_auth():
        return redirect(url_for('ui.login'))
    if not is_admin():
        flash("Permission denied. Administrator role required to create file shares.")
        return redirect(url_for('ui.fileshares'))
        
    name = request.form.get('name') or request.form.get('share_name')
    try:
        get_share_service().create_share(name)
        log_activity('file', 'CREATE_SHARE', name, 'SUCCESS')
        flash(f"File share '{name}' created.")
    except Exception as e:
        log_activity('file', 'CREATE_SHARE', name, 'FAILED', details=str(e))
        flash(f"Error creating file share: {e}")
    return redirect(url_for('ui.fileshares'))

@ui.route('/fileshares/delete', methods=['POST'])
def delete_share():
    if not require_auth():
        return redirect(url_for('ui.login'))
    if not is_admin():
        flash("Permission denied. Administrator role required to delete file shares.")
        return redirect(url_for('ui.fileshares'))
        
    name = request.form.get('share_name')
    try:
        get_share_service().delete_share(name)
        log_activity('file', 'DELETE_SHARE', name, 'SUCCESS')
        flash(f"File share '{name}' deleted.")
    except Exception as e:
        log_activity('file', 'DELETE_SHARE', name, 'FAILED', details=str(e))
        flash(f"Error deleting file share: {e}")
    return redirect(url_for('ui.fileshares'))

@ui.route('/fileshares/<share>', methods=['GET', 'POST'])
def list_files(share):
    if not require_auth():
        return redirect(url_for('ui.login'))

    path = request.args.get('path', '').strip('/')
    search_query = request.args.get('q', '').strip()
    from_date = request.args.get('from_date', '').strip()
    to_date = request.args.get('to_date', '').strip()

    from_dt = None
    to_dt = None
    if from_date:
        try:
            from_dt = datetime.strptime(from_date, "%Y-%m-%d").date()
        except Exception:
            pass
    if to_date:
        try:
            to_dt = datetime.strptime(to_date, "%Y-%m-%d").date()
        except Exception:
            pass

    try:
        page = max(1, int(request.args.get('page', 1)))
    except ValueError:
        page = 1
    try:
        limit = int(request.args.get('limit', 20))
        if limit not in [10, 15, 20, 50, 100]:
            limit = 20
    except ValueError:
        limit = 20

    sort_by = request.args.get('sort', 'name').strip().lower()
    order = request.args.get('order', 'asc').strip().lower()

    share_client = get_share_service().get_share_client(share)
    dir_client = share_client.get_directory_client(path)

    # Handle Directory Creation
    if request.method == 'POST' and 'dir_name' in request.form:
        if not has_write_permission():
            flash("Permission denied. Reader role is read-only.")
            return redirect(url_for('ui.list_files', share=share, path=path))
            
        new_dir = request.form.get('dir_name', '').strip()
        try:
            dir_client.create_subdirectory(new_dir)
            log_activity('file', 'CREATE_DIRECTORY', f"{share}/{path}/{new_dir}".strip('/'), 'SUCCESS')
            flash(f"Directory '{new_dir}' created.")
        except Exception as e:
            log_activity('file', 'CREATE_DIRECTORY', f"{share}/{path}/{new_dir}".strip('/'), 'FAILED', details=str(e))
            flash(f"Error creating directory: {e}")
        return redirect(url_for('ui.list_files', share=share, path=path))

    try:
        raw_items = list(dir_client.list_directories_and_files())
    except Exception as e:
        flash(f"Error loading files: {e}")
        raw_items = []

    filtered_items = []
    for item in raw_items:
        if search_query and search_query.lower() not in item['name'].lower():
            continue
        if (from_dt or to_dt) and not item.get('is_directory', False):
            lm = item.get('last_modified')
            if lm:
                f_date = lm.date()
                if from_dt and f_date < from_dt:
                    continue
                if to_dt and f_date > to_dt:
                    continue
        filtered_items.append(item)

    def sort_key(item):
        is_dir = 0 if item.get('is_directory', False) else 1
        if sort_by == 'size':
            return (is_dir, item.get('size', 0) or 0)
        elif sort_by == 'date':
            lm = item.get('last_modified')
            return (is_dir, lm.timestamp() if lm else 0)
        return (is_dir, item['name'].lower())

    reverse = (order == 'desc')
    filtered_items.sort(key=sort_key, reverse=reverse)

    total_items = len(filtered_items)
    total_pages = max(1, math.ceil(total_items / limit))
    if page > total_pages:
        page = total_pages

    start_idx = (page - 1) * limit
    end_idx = start_idx + limit
    paginated_items = filtered_items[start_idx:end_idx]

    # Breadcrumb parts
    parts = []
    accum = []
    if path:
        for p in path.split('/'):
            accum.append(p)
            parts.append({'name': p, 'path': '/'.join(accum)})

    tree = load_sidebar_tree()
    return render_template(
        'fileshares.html',
        items=paginated_items,
        share_name=share,
        path=path,
        path_parts=parts,
        search_query=search_query,
        page=page,
        total_pages=total_pages,
        limit=limit,
        total_items=total_items,
        sort_by=sort_by,
        order=order,
        from_date=from_date,
        to_date=to_date,
        sidebar_tree=tree,
        active_service='fileshares',
        active_item=share
    )

@ui.route('/fileshares/<share>/upload', methods=['POST'])
def upload_file(share):
    if not require_auth():
        return redirect(url_for('ui.login'))
    if not has_write_permission():
        flash("Permission denied. Reader role is read-only.")
        return redirect(url_for('ui.list_files', share=share))
        
    path = request.form.get('path', '').strip('/')
    files = request.files.getlist('files')

    if not files or not any(f.filename for f in files):
        flash("No files selected for upload.")
        return redirect(url_for('ui.list_files', share=share, path=path))

    share_client = get_share_service().get_share_client(share)
    dir_client = share_client.get_directory_client(path)

    uploaded = []
    failed = []
    for file in files:
        if not file or not file.filename:
            continue
        try:
            file_client = dir_client.get_file_client(file.filename)
            data = file.read()
            file_client.upload_file(data)
            uploaded.append(file.filename)
            log_activity('file', 'UPLOAD_FILE', f"{share}/{path}/{file.filename}".strip('/'), 'SUCCESS')
        except Exception as e:
            failed.append(f"{file.filename} ({e})")
            log_activity('file', 'UPLOAD_FILE', f"{share}/{path}/{file.filename}".strip('/'), 'FAILED', details=str(e))

    if uploaded:
        flash(f"Successfully uploaded {len(uploaded)} file(s).")
    if failed:
        flash(f"Failed to upload: {', '.join(failed)}")

    return redirect(url_for('ui.list_files', share=share, path=path))

@ui.route('/fileshares/<share>/download')
def download_file(share):
    if not require_auth():
        return redirect(url_for('ui.login'))
    path = request.args.get('path', '').strip('/')
    share_client = get_share_service().get_share_client(share)
    file_client = share_client.get_file_client(path)
    stream = io.BytesIO()
    file_client.download_file().readinto(stream)
    stream.seek(0)
    filename = secure_filename(path.split('/')[-1]) or 'download'
    return send_file(stream, download_name=filename, as_attachment=True)

@ui.route('/fileshares/<share>/file-content', methods=['GET'])
def get_file_content(share):
    if not require_auth():
        return jsonify({"success": False, "error": "Not authenticated"}), 401

    path = request.args.get('path', '').strip('/')
    if not path:
        return jsonify({"success": False, "error": "File path required"}), 400

    try:
        share_client = get_share_service().get_share_client(share)
        file_client = share_client.get_file_client(path)
        stream = io.BytesIO()
        file_client.download_file().readinto(stream)
        stream.seek(0)
        raw_bytes = stream.read()

        try:
            content = raw_bytes.decode('utf-8')
            return jsonify({"success": True, "content": content, "path": path})
        except UnicodeDecodeError:
            return jsonify({"success": False, "error": "Binary file cannot be edited in text editor."}), 400
    except Exception as e:
        return jsonify({"success": False, "error": str(e)}), 400

@ui.route('/fileshares/<share>/save-content', methods=['POST'])
def save_file_content(share):
    if not require_auth():
        return jsonify({"success": False, "error": "Not authenticated"}), 401
    if not has_write_permission():
        return jsonify({"success": False, "error": "Permission denied. Reader role is read-only."}), 403

    data = request.get_json(silent=True) or request.form
    path = data.get('path', '').strip('/')
    content = data.get('content')

    if not path or content is None:
        return jsonify({"success": False, "error": "Path and content are required"}), 400

    try:
        share_client = get_share_service().get_share_client(share)
        file_client = share_client.get_file_client(path)
        file_client.upload_file(content.encode('utf-8'))
        log_activity('file', 'EDIT_FILE_CONTENT', f"{share}/{path}", 'SUCCESS')
        return jsonify({"success": True, "message": "Saved successfully"})
    except Exception as e:
        log_activity('file', 'EDIT_FILE_CONTENT', f"{share}/{path}", 'FAILED', details=str(e))
        return jsonify({"success": False, "error": str(e)}), 400

@ui.route('/fileshares/<share>/rename-file', methods=['POST'])
def rename_file(share):
    if not require_auth():
        return jsonify({"success": False, "error": "Not authenticated"}), 401
    if not has_write_permission():
        return jsonify({"success": False, "error": "Permission denied. Reader role is read-only."}), 403

    data = request.get_json(silent=True) or request.form
    old_path = data.get('old_path', '').strip('/')
    new_name = data.get('new_name', '').strip()

    if not old_path or not new_name:
        return jsonify({"success": False, "error": "Path and new name are required"}), 400

    try:
        parent_dir = '/'.join(old_path.split('/')[:-1])
        new_path = f"{parent_dir}/{new_name}".strip('/')
        share_client = get_share_service().get_share_client(share)
        old_client = share_client.get_file_client(old_path)
        new_client = share_client.get_file_client(new_path)

        data = old_client.download_file().readall()
        new_client.upload_file(data)
        old_client.delete_file()
        log_activity('file', 'RENAME_FILE', f"{share}/{old_path} -> {new_path}", 'SUCCESS')
        return jsonify({"success": True, "message": f"Renamed to '{new_name}'."})
    except Exception as e:
        log_activity('file', 'RENAME_FILE', f"{share}/{old_path}", 'FAILED', details=str(e))
        return jsonify({"success": False, "error": str(e)}), 400

@ui.route('/fileshares/<share>/move-copy-file', methods=['POST'])
def move_copy_file(share):
    if not require_auth():
        return jsonify({"success": False, "error": "Not authenticated"}), 401
    if not has_write_permission():
        return jsonify({"success": False, "error": "Permission denied. Reader role is read-only."}), 403

    data = request.get_json(silent=True) or request.form
    src_path = data.get('src_path', '').strip('/')
    dest_share = data.get('dest_share', '').strip()
    dest_path = data.get('dest_path', '').strip('/')
    action_type = data.get('action_type', 'copy')

    if not src_path or not dest_share:
        return jsonify({"success": False, "error": "Source path and destination share are required"}), 400

    try:
        svc = get_share_service()
        src_client = svc.get_share_client(share).get_file_client(src_path)
        dest_client = svc.get_share_client(dest_share).get_file_client(dest_path or src_path.split('/')[-1])

        data = src_client.download_file().readall()
        dest_client.upload_file(data)
        if action_type == 'move':
            src_client.delete_file()
        log_activity('file', f"{action_type.upper()}_FILE", f"{share}/{src_path} -> {dest_share}/{dest_path}", 'SUCCESS')
        return jsonify({"success": True, "message": f"Successfully {'moved' if action_type == 'move' else 'copied'} file."})
    except Exception as e:
        log_activity('file', f"{action_type.upper()}_FILE", f"{share}/{src_path} -> {dest_share}/{dest_path}", 'FAILED', details=str(e))
        return jsonify({"success": False, "error": str(e)}), 400

@ui.route('/fileshares/<share>/delete', methods=['POST'])
def delete_file(share):
    if not require_auth():
        return redirect(url_for('ui.login'))
    if not has_write_permission():
        flash("Permission denied. Reader role is read-only.")
        return redirect(url_for('ui.list_files', share=share))
        
    path = request.form.get('path', '').strip('/')
    is_dir = request.form.get('is_directory') == 'true'
    current_path = request.form.get('current_path', '').strip('/')

    share_client = get_share_service().get_share_client(share)

    try:
        if is_dir:
            share_client.get_directory_client(path).delete_directory()
            log_activity('file', 'DELETE_DIRECTORY', f"{share}/{path}", 'SUCCESS')
            flash(f"Directory '{path}' deleted.")
        else:
            share_client.get_file_client(path).delete_file()
            log_activity('file', 'DELETE_FILE', f"{share}/{path}", 'SUCCESS')
            flash(f"File '{path}' deleted.")
    except Exception as e:
        log_activity('file', 'DELETE_FILE', f"{share}/{path}", 'FAILED', details=str(e))
        flash(f"Error deleting: {e}")

    return redirect(url_for('ui.list_files', share=share, path=current_path))

# -----------------------
# Queues Routes
# -----------------------
@ui.route('/queues')
def queues():
    if not require_auth():
        return redirect(url_for('ui.login'))

    try:
        queue_list = list(get_queue_service().list_queues())
    except Exception as e:
        flash(f"Error loading queues: {e}")
        queue_list = []

    tree = load_sidebar_tree()
    return render_template('queues.html', queues=queue_list, queue_name=None, sidebar_tree=tree, active_service='queues', active_item=None)

@ui.route('/queues/create', methods=['POST'])
def create_queue():
    if not require_auth():
        return redirect(url_for('ui.login'))
    if not is_admin():
        flash("Permission denied. Administrator role required to create queues.")
        return redirect(url_for('ui.queues'))
        
    name = request.form.get('name') or request.form.get('queue_name')
    try:
        get_queue_service().create_queue(name)
        log_activity('queue', 'CREATE_QUEUE', name, 'SUCCESS')
        flash(f"Queue '{name}' created.")
    except Exception as e:
        log_activity('queue', 'CREATE_QUEUE', name, 'FAILED', details=str(e))
        flash(f"Error creating queue: {e}")
    return redirect(url_for('ui.queues'))

@ui.route('/queues/delete', methods=['POST'])
def delete_queue():
    if not require_auth():
        return redirect(url_for('ui.login'))
    if not is_admin():
        flash("Permission denied. Administrator role required to delete queues.")
        return redirect(url_for('ui.queues'))
        
    name = request.form.get('queue_name')
    try:
        get_queue_service().delete_queue(name)
        log_activity('queue', 'DELETE_QUEUE', name, 'SUCCESS')
        flash(f"Queue '{name}' deleted.")
    except Exception as e:
        log_activity('queue', 'DELETE_QUEUE', name, 'FAILED', details=str(e))
        flash(f"Error deleting queue: {e}")
    return redirect(url_for('ui.queues'))

@ui.route('/queues/<queue>')
def view_queue(queue):
    if not require_auth():
        return redirect(url_for('ui.login'))

    search_query = request.args.get('q', '').strip()
    from_date = request.args.get('from_date', '').strip()
    to_date = request.args.get('to_date', '').strip()

    from_dt = None
    to_dt = None
    if from_date:
        try:
            from_dt = datetime.strptime(from_date, "%Y-%m-%d").date()
        except Exception:
            pass
    if to_date:
        try:
            to_dt = datetime.strptime(to_date, "%Y-%m-%d").date()
        except Exception:
            pass

    try:
        page = max(1, int(request.args.get('page', 1)))
    except ValueError:
        page = 1
    try:
        limit = int(request.args.get('limit', 20))
        if limit not in [10, 15, 20, 50, 100]:
            limit = 20
    except ValueError:
        limit = 20

    sort_by = request.args.get('sort', 'inserted_on').strip().lower()
    order = request.args.get('order', 'desc').strip().lower()

    queue_client = get_queue_service().get_queue_client(queue)

    try:
        raw_msgs = list(queue_client.peek_messages(max_messages=32))
    except Exception as e:
        flash(f"Error reading queue: {e}")
        raw_msgs = []

    filtered_msgs = []
    for msg in raw_msgs:
        if search_query and search_query.lower() not in (msg.content or '').lower() and search_query.lower() not in msg.id.lower():
            continue
        if from_dt or to_dt:
            lm = getattr(msg, 'insertion_time', None)
            if lm:
                m_date = lm.date()
                if from_dt and m_date < from_dt:
                    continue
                if to_dt and m_date > to_dt:
                    continue
        filtered_msgs.append(msg)

    def sort_key(msg):
        if sort_by == 'inserted_on':
            return msg.insertion_time.timestamp() if getattr(msg, 'insertion_time', None) else 0
        elif sort_by == 'expires_on':
            return msg.expiration_time.timestamp() if getattr(msg, 'expiration_time', None) else 0
        elif sort_by == 'dequeue_count':
            return getattr(msg, 'dequeue_count', 0) or 0
        return msg.id

    reverse = (order == 'desc')
    filtered_msgs.sort(key=sort_key, reverse=reverse)

    total_items = len(filtered_msgs)
    total_pages = max(1, math.ceil(total_items / limit))
    if page > total_pages:
        page = total_pages

    start_idx = (page - 1) * limit
    end_idx = start_idx + limit
    paginated_msgs = filtered_msgs[start_idx:end_idx]

    tree = load_sidebar_tree()
    return render_template(
        'queues.html',
        messages=paginated_msgs,
        queue_name=queue,
        search_query=search_query,
        page=page,
        total_pages=total_pages,
        limit=limit,
        total_items=total_items,
        sort_by=sort_by,
        order=order,
        from_date=from_date,
        to_date=to_date,
        sidebar_tree=tree,
        active_service='queues',
        active_item=queue
    )

@ui.route('/queues/<queue>/enqueue', methods=['POST'])
def enqueue_message(queue):
    if not require_auth():
        return redirect(url_for('ui.login'))
    if not has_write_permission():
        flash("Permission denied. Reader role is read-only.")
        return redirect(url_for('ui.view_queue', queue=queue))
        
    content = request.form.get('content', '').strip()
    ttl = int(request.form.get('ttl', 604800) or 604800)
    visibility_timeout = int(request.form.get('visibility_timeout', 0) or 0)

    try:
        queue_client = get_queue_service().get_queue_client(queue)
        queue_client.send_message(content, time_to_live=ttl, visibility_timeout=visibility_timeout)
        log_activity('queue', 'ENQUEUE_MESSAGE', queue, 'SUCCESS', details=f"Size: {len(content)} chars, TTL: {ttl}s")
        flash("Message enqueued successfully.")
    except Exception as e:
        log_activity('queue', 'ENQUEUE_MESSAGE', queue, 'FAILED', details=str(e))
        flash(f"Error enqueuing message: {e}")

    return redirect(url_for('ui.view_queue', queue=queue))

@ui.route('/queues/<queue>/message-content', methods=['GET'])
def get_queue_message_content(queue):
    if not require_auth():
        return jsonify({'success': False, 'error': 'Unauthorized'}), 401
    msg_id = request.args.get('id')
    if not msg_id:
        return jsonify({'success': False, 'error': 'Message ID is required'}), 400
    try:
        queue_client = get_queue_service().get_queue_client(queue)
        messages = list(queue_client.peek_messages(max_messages=32))
        for msg in messages:
            if msg.id == msg_id:
                return jsonify({
                    'success': True,
                    'id': msg.id,
                    'content': msg.content,
                    'insertion_time': msg.insertion_time.isoformat() if hasattr(msg, 'insertion_time') and msg.insertion_time else '',
                    'expiration_time': msg.expiration_time.isoformat() if hasattr(msg, 'expiration_time') and msg.expiration_time else '',
                    'dequeue_count': getattr(msg, 'dequeue_count', 0)
                })
        return jsonify({'success': False, 'error': 'Message not found'}), 404
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 400

@ui.route('/queues/<queue>/update-message', methods=['POST'])
def update_queue_message(queue):
    if not require_auth():
        return jsonify({'success': False, 'error': 'Unauthorized'}), 401
    if not has_write_permission():
        return jsonify({'success': False, 'error': 'Permission denied. Reader role is read-only.'}), 403
        
    data = request.get_json(silent=True) or request.form
    msg_id = data.get('id')
    content = data.get('content')
    if not msg_id or content is None:
        return jsonify({'success': False, 'error': 'Message ID and Content are required'}), 400
    try:
        queue_client = get_queue_service().get_queue_client(queue)
        msgs = list(queue_client.receive_messages(messages_per_page=32, visibility_timeout=30))
        target_msg = None
        for m in msgs:
            if m.id == msg_id:
                target_msg = m
                break
        if not target_msg:
            return jsonify({'success': False, 'error': 'Message could not be leased or does not exist.'}), 404
            
        queue_client.update_message(target_msg.id, target_msg.pop_receipt, content=content, visibility_timeout=0)
        log_activity('queue', 'UPDATE_MESSAGE', queue, 'SUCCESS', details=f"Message ID: {msg_id}")
        return jsonify({'success': True, 'message': 'Message updated successfully.'})
    except Exception as e:
        log_activity('queue', 'UPDATE_MESSAGE', queue, 'FAILED', details=str(e))
        return jsonify({'success': False, 'error': str(e)}), 400

@ui.route('/queues/<queue>/dequeue-single', methods=['POST'])
def dequeue_single_message(queue):
    if not require_auth():
        return jsonify({'success': False, 'error': 'Unauthorized'}), 401
    if not has_write_permission():
        return jsonify({'success': False, 'error': 'Permission denied. Reader role is read-only.'}), 403
        
    data = request.get_json(silent=True) or request.form
    msg_id = data.get('id')
    if not msg_id:
        return jsonify({'success': False, 'error': 'Message ID required'}), 400
    try:
        queue_client = get_queue_service().get_queue_client(queue)
        msgs = list(queue_client.receive_messages(messages_per_page=32, visibility_timeout=30))
        target_msg = None
        for m in msgs:
            if m.id == msg_id:
                target_msg = m
                break
        if not target_msg:
            return jsonify({'success': False, 'error': 'Message could not be leased for deletion.'}), 404
        queue_client.delete_message(target_msg.id, target_msg.pop_receipt)
        log_activity('queue', 'DEQUEUE_SINGLE', queue, 'SUCCESS', details=f"Message ID: {msg_id}")
        return jsonify({'success': True, 'message': 'Message dequeued and deleted successfully.'})
    except Exception as e:
        log_activity('queue', 'DEQUEUE_SINGLE', queue, 'FAILED', details=str(e))
        return jsonify({'success': False, 'error': str(e)}), 400

@ui.route('/queues/<queue>/dequeue-multiple', methods=['POST'])
def dequeue_multiple_messages(queue):
    if not require_auth():
        return jsonify({'success': False, 'error': 'Unauthorized'}), 401
    if not has_write_permission():
        return jsonify({'success': False, 'error': 'Permission denied. Reader role is read-only.'}), 403
        
    data = request.get_json(silent=True) or request.form
    msg_ids = data.get('ids', [])
    if not msg_ids:
        return jsonify({'success': False, 'error': 'No message IDs specified'}), 400
    try:
        queue_client = get_queue_service().get_queue_client(queue)
        msgs = list(queue_client.receive_messages(messages_per_page=32, visibility_timeout=30))
        deleted_count = 0
        for m in msgs:
            if m.id in msg_ids:
                queue_client.delete_message(m.id, m.pop_receipt)
                deleted_count += 1
        log_activity('queue', 'DEQUEUE_MULTIPLE', queue, 'SUCCESS', details=f"Dequeued {deleted_count} messages")
        return jsonify({'success': True, 'deleted_count': deleted_count, 'message': f"Dequeued {deleted_count} messages successfully."})
    except Exception as e:
        log_activity('queue', 'DEQUEUE_MULTIPLE', queue, 'FAILED', details=str(e))
        return jsonify({'success': False, 'error': str(e)}), 400

@ui.route('/queues/<queue>/send-to-queue', methods=['POST'])
def send_to_queue(queue):
    if not require_auth():
        return jsonify({'success': False, 'error': 'Unauthorized'}), 401
    if not has_write_permission():
        return jsonify({'success': False, 'error': 'Permission denied. Reader role is read-only.'}), 403
        
    data = request.get_json(silent=True) or request.form
    msg_id = data.get('id')
    dest_queue = data.get('dest_queue')
    action_type = data.get('action_type', 'copy') # 'move' or 'copy'
    if not msg_id or not dest_queue:
        return jsonify({'success': False, 'error': 'Message ID and Destination Queue are required'}), 400
    try:
        svc = get_queue_service()
        src_client = svc.get_queue_client(queue)
        dest_client = svc.get_queue_client(dest_queue)
        
        msgs = list(src_client.receive_messages(messages_per_page=32, visibility_timeout=30))
        target_msg = None
        for m in msgs:
            if m.id == msg_id:
                target_msg = m
                break
        if not target_msg:
            return jsonify({'success': False, 'error': 'Message could not be leased or found.'}), 404
            
        dest_client.send_message(target_msg.content)
        if action_type == 'move':
            src_client.delete_message(target_msg.id, target_msg.pop_receipt)
            
        log_activity('queue', f"SEND_TO_QUEUE_{action_type.upper()}", f"{queue} -> {dest_queue}", 'SUCCESS', details=f"Message ID: {msg_id}")
        return jsonify({'success': True, 'message': f"Message {'moved' if action_type == 'move' else 'copied'} to queue '{dest_queue}'."})
    except Exception as e:
        log_activity('queue', f"SEND_TO_QUEUE_{action_type.upper()}", f"{queue} -> {dest_queue}", 'FAILED', details=str(e))
        return jsonify({'success': False, 'error': str(e)}), 400

@ui.route('/queues/<queue>/dequeue', methods=['POST'])
def dequeue_message(queue):
    if not require_auth():
        return redirect(url_for('ui.login'))
    if not has_write_permission():
        flash("Permission denied. Reader role is read-only.")
        return redirect(url_for('ui.view_queue', queue=queue))
        
    try:
        queue_client = get_queue_service().get_queue_client(queue)
        messages = queue_client.receive_messages(messages_per_page=1, visibility_timeout=30)
        count = 0
        for msg in messages:
            queue_client.delete_message(msg)
            count += 1
        log_activity('queue', 'DEQUEUE_MESSAGE', queue, 'SUCCESS')
        flash(f"Dequeued {count} message.")
    except Exception as e:
        log_activity('queue', 'DEQUEUE_MESSAGE', queue, 'FAILED', details=str(e))
        flash(f"Error dequeuing: {e}")
    return redirect(url_for('ui.view_queue', queue=queue))

@ui.route('/queues/<queue>/dequeue-all', methods=['POST'])
def dequeue_all_messages(queue):
    if not require_auth():
        return redirect(url_for('ui.login'))
    if not has_write_permission():
        flash("Permission denied. Reader role is read-only.")
        return redirect(url_for('ui.view_queue', queue=queue))

    try:
        queue_client = get_queue_service().get_queue_client(queue)
        queue_client.clear_messages()
        log_activity('queue', 'DEQUEUE_ALL', queue, 'SUCCESS')
        flash(f"Cleared all messages from queue '{queue}'.")
    except Exception as e:
        log_activity('queue', 'DEQUEUE_ALL', queue, 'FAILED', details=str(e))
        flash(f"Error clearing queue: {e}")
    return redirect(url_for('ui.view_queue', queue=queue))

# -----------------------
# Tables Routes
# -----------------------
@ui.route('/tables')
def list_tables():
    if not require_auth():
        return redirect(url_for('ui.login'))

    try:
        table_svc = get_table_service()
        raw_tables = list(table_svc.list_tables())
        table_list = [t.name if hasattr(t, "name") else str(t) for t in raw_tables]
        # Privacy: Filter out USER_TABLE and LOGS_TABLE for non-admins
        if not is_admin():
            table_list = [t for t in table_list if t.lower() not in [USER_TABLE.lower(), LOGS_TABLE.lower()]]
    except Exception as e:
        flash(f"Error loading tables: {e}")
        table_list = []

    tree = load_sidebar_tree()
    return render_template('tables.html', tables=table_list, table_name=None, sidebar_tree=tree, active_service='tables', active_item=None)

@ui.route('/tables/create', methods=['POST'])
def create_table():
    if not require_auth():
        return redirect(url_for('ui.login'))
    if not is_admin():
        flash("Permission denied. Administrator role required to create tables.")
        return redirect(url_for('ui.list_tables'))
        
    name = request.form.get('name') or request.form.get('table_name')
    try:
        get_table_service().create_table(name)
        log_activity('table', 'CREATE_TABLE', name, 'SUCCESS')
        flash(f"Table '{name}' created.")
    except Exception as e:
        log_activity('table', 'CREATE_TABLE', name, 'FAILED', details=str(e))
        flash(f"Error creating table: {e}")
    return redirect(url_for('ui.list_tables'))

@ui.route('/tables/delete', methods=['POST'])
def delete_table():
    if not require_auth():
        return redirect(url_for('ui.login'))
    if not is_admin():
        flash("Permission denied. Administrator role required to delete tables.")
        return redirect(url_for('ui.list_tables'))
        
    name = request.form.get('table_name')
    if name.lower() in [USER_TABLE.lower(), LOGS_TABLE.lower()]:
        flash("System tables cannot be deleted.")
        return redirect(url_for('ui.list_tables'))
        
    try:
        get_table_service().delete_table(name)
        log_activity('table', 'DELETE_TABLE', name, 'SUCCESS')
        flash(f"Table '{name}' deleted.")
    except Exception as e:
        log_activity('table', 'DELETE_TABLE', name, 'FAILED', details=str(e))
        flash(f"Error deleting table: {e}")
    return redirect(url_for('ui.list_tables'))

@ui.route('/tables/<table_name>')
def view_table(table_name):
    if not require_auth():
        return redirect(url_for('ui.login'))

    # Security & Privacy: Restrict system tables to Admin only
    if table_name.lower() in [USER_TABLE.lower(), LOGS_TABLE.lower()] and not is_admin():
        flash("Access to system tables is restricted to administrators.")
        return redirect(url_for('ui.list_tables'))

    filter_expr = request.args.get('filter', '').strip()
    search_query = request.args.get('q', '').strip()
    from_date = request.args.get('from_date', '').strip()
    to_date = request.args.get('to_date', '').strip()

    from_dt = None
    to_dt = None
    if from_date:
        try:
            from_dt = datetime.strptime(from_date, "%Y-%m-%d").date()
        except Exception:
            pass
    if to_date:
        try:
            to_dt = datetime.strptime(to_date, "%Y-%m-%d").date()
        except Exception:
            pass

    try:
        page = max(1, int(request.args.get('page', 1)))
    except ValueError:
        page = 1
    try:
        limit = int(request.args.get('limit', 20))
        if limit not in [10, 15, 20, 50, 100]:
            limit = 20
    except ValueError:
        limit = 20

    sort_by = request.args.get('sort', 'PartitionKey').strip()
    order = request.args.get('order', 'asc').strip().lower()

    table_client = get_table_service().get_table_client(table_name)

    try:
        if filter_expr:
            entities = list(table_client.query_entities(filter_expr))
        else:
            entities = list(table_client.list_entities())
    except Exception as e:
        flash(f"Error querying table: {e}")
        entities = []

    # Dynamic Column Discovery
    all_keys = set()
    dict_entities = []
    for e in entities:
        d = dict(e)
        # Omit internal OData metadata
        d.pop('odata.etag', None)
        d.pop('etag', None)
        all_keys.update(d.keys())
        dict_entities.append(d)

    # Place PartitionKey and RowKey and Timestamp at front
    cols = []
    if 'PartitionKey' in all_keys:
        cols.append('PartitionKey')
    if 'RowKey' in all_keys:
        cols.append('RowKey')
    if 'Timestamp' in all_keys:
        cols.append('Timestamp')
    for k in sorted(all_keys):
        if k not in cols:
            cols.append(k)

    # In-memory search filter & date filter
    filtered_entities = []
    for d in dict_entities:
        if search_query:
            match = False
            for v in d.values():
                if search_query.lower() in str(v).lower():
                    match = True
                    break
            if not match:
                continue
        if from_dt or to_dt:
            ts = d.get('Timestamp')
            if ts:
                try:
                    if hasattr(ts, 'date'):
                        t_date = ts.date()
                    else:
                        t_date = datetime.fromisoformat(str(ts).replace('Z', '+00:00')).date()
                    if from_dt and t_date < from_dt:
                        continue
                    if to_dt and t_date > to_dt:
                        continue
                except Exception:
                    pass
        filtered_entities.append(d)

    # Sorting
    def sort_key(d):
        val = d.get(sort_by, '')
        if val is None:
            return ''
        return str(val).lower()

    reverse = (order == 'desc')
    filtered_entities.sort(key=sort_key, reverse=reverse)

    total_items = len(filtered_entities)
    total_pages = max(1, math.ceil(total_items / limit))
    if page > total_pages:
        page = total_pages

    start_idx = (page - 1) * limit
    end_idx = start_idx + limit
    paginated_entities = filtered_entities[start_idx:end_idx]

    tree = load_sidebar_tree()
    return render_template(
        'tables.html',
        entities=paginated_entities,
        columns=cols,
        table_name=table_name,
        filter_expr=filter_expr,
        search_query=search_query,
        page=page,
        total_pages=total_pages,
        limit=limit,
        total_items=total_items,
        sort_by=sort_by,
        order=order,
        from_date=from_date,
        to_date=to_date,
        sidebar_tree=tree,
        active_service='tables',
        active_item=table_name
    )

@ui.route('/tables/<table_name>/add', methods=['POST'])
def add_entity(table_name):
    if not require_auth():
        return redirect(url_for('ui.login'))
    if not has_write_permission():
        flash("Permission denied. Reader role is read-only.")
        return redirect(url_for('ui.view_table', table_name=table_name))
    if table_name.lower() in [USER_TABLE.lower(), LOGS_TABLE.lower()] and not is_admin():
        flash("Access to system tables is restricted to administrators.")
        return redirect(url_for('ui.list_tables'))

    pk = request.form.get('pk')
    rk = request.form.get('rk')
    json_data = request.form.get('entity_json', '{}')
    try:
        custom_props = json.loads(json_data)
        entity = {'PartitionKey': pk, 'RowKey': rk}
        entity.update(custom_props)
        get_table_service().get_table_client(table_name).create_entity(entity)
        log_activity('table', 'INSERT_ENTITY', f"{table_name} (PK={pk}, RK={rk})", 'SUCCESS')
        flash("Entity added.")
    except Exception as e:
        log_activity('table', 'INSERT_ENTITY', f"{table_name} (PK={pk}, RK={rk})", 'FAILED', details=str(e))
        flash(f"Error adding entity: {e}")
    return redirect(url_for('ui.view_table', table_name=table_name))

@ui.route('/tables/<table_name>/entity-content', methods=['GET'])
def get_table_entity_content(table_name):
    if not require_auth():
        return jsonify({'success': False, 'error': 'Unauthorized'}), 401
    if table_name.lower() in [USER_TABLE.lower(), LOGS_TABLE.lower()] and not is_admin():
        return jsonify({'success': False, 'error': 'Access to system tables is restricted to administrators.'}), 403

    pk = request.args.get('pk')
    rk = request.args.get('rk')
    if not pk or not rk:
        return jsonify({'success': False, 'error': 'PartitionKey and RowKey are required.'}), 400
        
    try:
        client = get_table_service().get_table_client(table_name)
        entity = client.get_entity(partition_key=pk, row_key=rk)
        dict_entity = {k: v for k, v in entity.items() if k not in ['odata.etag', 'etag']}
        if 'Timestamp' in dict_entity and hasattr(dict_entity['Timestamp'], 'isoformat'):
            dict_entity['Timestamp'] = dict_entity['Timestamp'].isoformat()
        return jsonify({'success': True, 'entity': dict_entity})
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 400

@ui.route('/tables/<table_name>/update-entity', methods=['POST'])
def update_table_entity(table_name):
    if not require_auth():
        return jsonify({'success': False, 'error': 'Unauthorized'}), 401
    if not has_write_permission():
        return jsonify({'success': False, 'error': 'Permission denied. Reader role is read-only.'}), 403
    if table_name.lower() in [USER_TABLE.lower(), LOGS_TABLE.lower()] and not is_admin():
        return jsonify({'success': False, 'error': 'Access to system tables is restricted to administrators.'}), 403

    data = request.get_json(silent=True) or request.form
    pk = data.get('pk')
    rk = data.get('rk')
    entity_data = data.get('entity')

    if not pk or not rk or entity_data is None:
        return jsonify({'success': False, 'error': 'PartitionKey, RowKey and entity data are required.'}), 400
        
    try:
        if isinstance(entity_data, str):
            entity_data = json.loads(entity_data)
            
        entity_data['PartitionKey'] = pk
        entity_data['RowKey'] = rk
        entity_data.pop('Timestamp', None)
        entity_data.pop('odata.etag', None)
        entity_data.pop('etag', None)
        
        client = get_table_service().get_table_client(table_name)
        client.upsert_entity(entity=entity_data, mode=UpdateMode.REPLACE)
        log_activity('table', 'UPDATE_ENTITY', f"{table_name} (PK={pk}, RK={rk})", 'SUCCESS')
        return jsonify({'success': True, 'message': 'Entity updated successfully.'})
    except Exception as e:
        log_activity('table', 'UPDATE_ENTITY', f"{table_name} (PK={pk}, RK={rk})", 'FAILED', details=str(e))
        return jsonify({'success': False, 'error': str(e)}), 400

@ui.route('/tables/<table_name>/clone-entity', methods=['POST'])
def clone_table_entity(table_name):
    if not require_auth():
        return jsonify({'success': False, 'error': 'Unauthorized'}), 401
    if not has_write_permission():
        return jsonify({'success': False, 'error': 'Permission denied. Reader role is read-only.'}), 403
    if table_name.lower() in [USER_TABLE.lower(), LOGS_TABLE.lower()] and not is_admin():
        return jsonify({'success': False, 'error': 'Access to system tables is restricted to administrators.'}), 403

    data = request.get_json(silent=True) or request.form
    src_pk = data.get('source_pk')
    src_rk = data.get('source_rk')
    new_pk = data.get('new_pk')
    new_rk = data.get('new_rk')
    
    if not src_pk or not src_rk or not new_pk or not new_rk:
        return jsonify({'success': False, 'error': 'Source and target PartitionKey and RowKey are required.'}), 400
        
    try:
        client = get_table_service().get_table_client(table_name)
        source_ent = client.get_entity(partition_key=src_pk, row_key=src_rk)
        new_ent = {k: v for k, v in source_ent.items() if k not in ['odata.etag', 'etag', 'Timestamp']}
        new_ent['PartitionKey'] = new_pk
        new_ent['RowKey'] = new_rk
        
        client.create_entity(new_ent)
        log_activity('table', 'CLONE_ENTITY', f"{table_name} (PK={src_pk}, RK={src_rk}) -> (PK={new_pk}, RK={new_rk})", 'SUCCESS')
        return jsonify({'success': True, 'message': f"Cloned entity to PartitionKey='{new_pk}', RowKey='{new_rk}'."})
    except Exception as e:
        log_activity('table', 'CLONE_ENTITY', f"{table_name} (PK={src_pk}, RK={src_rk}) -> (PK={new_pk}, RK={new_rk})", 'FAILED', details=str(e))
        return jsonify({'success': False, 'error': str(e)}), 400

@ui.route('/tables/<table_name>/copy-entity', methods=['POST'])
def copy_table_entity(table_name):
    if not require_auth():
        return jsonify({'success': False, 'error': 'Unauthorized'}), 401
    if not has_write_permission():
        return jsonify({'success': False, 'error': 'Permission denied. Reader role is read-only.'}), 403
    if table_name.lower() in [USER_TABLE.lower(), LOGS_TABLE.lower()] and not is_admin():
        return jsonify({'success': False, 'error': 'Access to system tables is restricted to administrators.'}), 403

    data = request.get_json(silent=True) or request.form
    pk = data.get('pk')
    rk = data.get('rk')
    dest_table = data.get('dest_table')
    action_type = data.get('action_type', 'copy')
    
    if not pk or not rk or not dest_table:
        return jsonify({'success': False, 'error': 'PartitionKey, RowKey and destination table name are required.'}), 400
        
    try:
        svc = get_table_service()
        src_client = svc.get_table_client(table_name)
        dest_client = svc.get_table_client(dest_table)
        
        source_ent = src_client.get_entity(partition_key=pk, row_key=rk)
        new_ent = {k: v for k, v in source_ent.items() if k not in ['odata.etag', 'etag', 'Timestamp']}
        new_ent['PartitionKey'] = pk
        new_ent['RowKey'] = rk
        
        dest_client.upsert_entity(entity=new_ent, mode=UpdateMode.REPLACE)
        if action_type == 'move':
            src_client.delete_entity(partition_key=pk, row_key=rk)
            
        log_activity('table', f"{action_type.upper()}_ENTITY", f"{table_name} -> {dest_table} (PK={pk}, RK={rk})", 'SUCCESS')
        return jsonify({'success': True, 'message': f"Entity successfully {'moved' if action_type == 'move' else 'copied'} to table '{dest_table}'."})
    except Exception as e:
        log_activity('table', f"{action_type.upper()}_ENTITY", f"{table_name} -> {dest_table} (PK={pk}, RK={rk})", 'FAILED', details=str(e))
        return jsonify({"success": False, "error": str(e)}), 400

@ui.route('/tables/<table_name>/delete', methods=['POST'])
def delete_entity(table_name):
    if not require_auth():
        return redirect(url_for('ui.login'))
    if not has_write_permission():
        flash("Permission denied. Reader role is read-only.")
        return redirect(url_for('ui.view_table', table_name=table_name))
    if table_name.lower() in [USER_TABLE.lower(), LOGS_TABLE.lower()] and not is_admin():
        flash("Access to system tables is restricted to administrators.")
        return redirect(url_for('ui.list_tables'))

    pk = request.form.get('pk')
    rk = request.form.get('rk')
    try:
        get_table_service().get_table_client(table_name).delete_entity(pk, rk)
        log_activity('table', 'DELETE_ENTITY', f"{table_name} (PK={pk}, RK={rk})", 'SUCCESS')
        flash("Entity deleted.")
    except Exception as e:
        log_activity('table', 'DELETE_ENTITY', f"{table_name} (PK={pk}, RK={rk})", 'FAILED', details=str(e))
        flash(f"Delete failed: {e}")
    return redirect(url_for('ui.view_table', table_name=table_name))

# Register Blueprint
app.register_blueprint(ui)

# -----------------------
# Root Redirect
# -----------------------
@app.route("/")
def index_redirect():
    return redirect("/storage-ui/")

# -----------------------
# Run
# -----------------------
if __name__ == '__main__':
    app.run(host='0.0.0.0', port=int(os.environ.get('PORT', 8080)), debug=True)
